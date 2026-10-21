"""Prevention & Healing Records Workspace test suite.

Covers:
  - Enhancement Programs CRUD & decimal expense calculations
  - Family Support Services CRUD & legacy RedMane ID handling
  - Other Nations Support CRUD
  - Inactive Staff Onboarding workflow (pending -> verified email -> activated & historical linking)
  - Canonical Client linking & reconciliation queue (10-digit Person ID)
  - Safe Excel Import Wizard (Analyze -> Preview -> Approval -> Execution)
  - Annual Summary reporting & financial aggregations
  - Role-based security & IT Admin boundary enforcement
"""

from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.org_ops import Employee
from app.models.person import Person
from app.models.prevention_healing import (
    EnhancementProgram,
    FamilySupportRecord,
    OtherNationsSupport,
    PHClientRef,
    PHImportBatch,
    PendingStaffProfile,
)


@pytest.mark.asyncio
async def test_enhancement_program_crud_and_expense_totals(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    headers = executive_director_user["headers"]

    payload = {
        "program": "Family Wellness & Culture",
        "program_activity": "Beading & Traditional Parenting Workshop",
        "record_date": "2026-05-15",
        "date_provided": "2026-05-15",
        "city": "Regina",
        "province": "SK",
        "partners": "Cowessess Knowledge Keepers",
        "families_in_province": 12,
        "families_out_of_province": 2,
        "participants_children_0_12": 15,
        "participants_youth_13_17": 8,
        "participants_post_majority_18_26": 4,
        "participants_parents_caregivers": 14,
        "participants_family_members": 6,
        "participants_others": 2,
        "account_code": "41000",
        "subcode": "213",
        "expense_gift_card": "250.00",
        "expense_purchase_order": "500.00",
        "expense_eft": "150.00",
        "expense_cheque": "0.00",
        "expense_credit_card": "125.50",
        "expense_courtesy_card": "0.00",
        "comments": "High engagement, traditional teachings well received.",
        "fiscal_year": "2026-2027",
    }

    # 1. Create
    res = await client.post("/api/v1/ph/enhancement-programs", json=payload, headers=headers)
    assert res.status_code == 201, res.text
    data = res.json()
    rec_id = data["id"]
    assert data["total_families"] == 14  # 12 + 2
    assert data["total_participants"] == 49  # 15 + 8 + 4 + 14 + 6 + 2
    assert data["total_cost"] == 1025.50  # 250 + 500 + 150 + 125.50

    # 2. Get
    get_res = await client.get(f"/api/v1/ph/enhancement-programs/{rec_id}", headers=headers)
    assert get_res.status_code == 200
    assert get_res.json()["program"] == "Family Wellness & Culture"

    # 3. List
    list_res = await client.get("/api/v1/ph/enhancement-programs?fiscal_year=2026-2027", headers=headers)
    assert list_res.status_code == 200
    assert len(list_res.json()) >= 1

    # 4. Update
    payload["program_activity"] = "Updated Traditional Parenting Workshop"
    up_res = await client.put(f"/api/v1/ph/enhancement-programs/{rec_id}", json=payload, headers=headers)
    assert up_res.status_code == 200
    assert up_res.json()["program_activity"] == "Updated Traditional Parenting Workshop"

    # 5. Delete
    del_res = await client.delete(f"/api/v1/ph/enhancement-programs/{rec_id}", headers=headers)
    assert del_res.status_code == 204


@pytest.mark.asyncio
async def test_family_support_crud_and_redmane_id(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    headers = executive_director_user["headers"]

    payload = {
        "program": "Family Support Services",
        "support_type": "Emergency Food & Utility",
        "service_area": "Regina Urban",
        "record_date": "2026-06-10",
        "date_provided": "2026-06-10",
        "client_ref_legacy_id": "RM-99482",
        "city": "Regina",
        "province": "SK",
        "participants_children_0_12": 3,
        "participants_youth_13_17": 1,
        "participants_post_majority_18_26": 0,
        "participants_adults_18_plus": 2,
        "account_code": "41000",
        "subcode": "110",
        "project_code": "215",
        "expense_gift_card": "200.00",
        "expense_purchase_order": "350.00",
        "expense_eft": "0.00",
        "expense_cheque": "0.00",
        "expense_credit_card": "0.00",
        "expense_petty_cash": "50.00",
        "expense_courtesy_card": "0.00",
        "comments": "Groceries voucher and emergency power support.",
        "fiscal_year": "2026-2027",
    }

    # 1. Create
    res = await client.post("/api/v1/ph/family-support", json=payload, headers=headers)
    assert res.status_code == 201, res.text
    data = res.json()
    rec_id = data["id"]
    assert data["client_ref_legacy_id"] == "RM-99482"
    assert data["total_cost"] == 600.00  # 200 + 350 + 50
    assert data["client_match_status"] == "UNRESOLVED"

    # 2. List
    list_res = await client.get("/api/v1/ph/family-support?fiscal_year=2026-2027", headers=headers)
    assert list_res.status_code == 200
    assert len(list_res.json()) >= 1

    # 3. Delete
    del_res = await client.delete(f"/api/v1/ph/family-support/{rec_id}", headers=headers)
    assert del_res.status_code == 204


@pytest.mark.asyncio
async def test_other_nations_support_crud(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    headers = executive_director_user["headers"]

    payload = {
        "worker_name_raw": "Tammy Breen",
        "record_date": "2026-07-27",
        "family_reference_raw": "Delorme Family",
        "service_type": "Utility Arrears",
        "cost": "2812.82",
        "on_reserve": False,
        "cowessess_connection": "Married to band member",
        "other_community": "Zagimay",
        "historical_reference_number": "3640201301",
        "comments": "Off reserve support coordination.",
        "fiscal_year": "2026-2027",
    }

    # 1. Create
    res = await client.post("/api/v1/ph/other-nations", json=payload, headers=headers)
    assert res.status_code == 201, res.text
    data = res.json()
    rec_id = data["id"]
    assert data["worker_name_raw"] == "Tammy Breen"
    assert data["cost"] == 2812.82
    assert data["on_reserve"] is False

    # 2. Get
    get_res = await client.get(f"/api/v1/ph/other-nations/{rec_id}", headers=headers)
    assert get_res.status_code == 200
    assert get_res.json()["other_community"] == "Zagimay"

    # 3. Update
    payload["comments"] = "Updated confirmation received from band administration."
    up_res = await client.put(f"/api/v1/ph/other-nations/{rec_id}", json=payload, headers=headers)
    assert up_res.status_code == 200
    assert up_res.json()["comments"] == "Updated confirmation received from band administration."


@pytest.mark.asyncio
async def test_staff_onboarding_inactive_and_activation_workflow(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    """Verifies:

    - Pending profile is created without email
    - Profile remains PENDING and cannot be activated without verified email
    - Providing verified email and activating creates an active Employee profile
    - Historical records referencing pending staff profile are automatically linked to Employee
    """
    headers = executive_director_user["headers"]

    # 1. Create pending staff profile from Excel without email
    staff_payload = {
        "first_name": "Valerie",
        "last_name": "Standingready",
        "display_name_from_excel": "Valerie Standingready",
        "position": "Prevention Worker",
        "department": "Prevention & Healing",
        "email": None,
        "phone": "306-555-0199",
        "notes": "Identified in historical 2026 Prevention tracker",
        "onboarding_status": "PENDING",
    }
    create_res = await client.post("/api/v1/ph/pending-staff", json=staff_payload, headers=headers)
    assert create_res.status_code == 201
    staff_data = create_res.json()
    profile_id = staff_data["id"]
    assert staff_data["onboarding_status"] == "PENDING"
    assert staff_data["employee_id"] is None

    # 2. Create historical EnhancementProgram referencing this pending staff profile
    ep = EnhancementProgram(
        pending_staff_id=uuid.UUID(profile_id),
        employee_name_raw="Valerie Standingready",
        program="Youth Cultural Camp",
        record_date=date(2026, 6, 1),
        fiscal_year="2026-2027",
    )
    db_session.add(ep)
    await db_session.commit()
    await db_session.refresh(ep)
    assert ep.employee_id is None

    # 3. Attempt to activate WITHOUT email -> must fail with 400
    fail_act = await client.post(f"/api/v1/ph/pending-staff/{profile_id}/activate", json={}, headers=headers)
    assert fail_act.status_code == 400
    assert "email" in fail_act.text.lower()

    # 4. Activate WITH verified email
    act_res = await client.post(
        f"/api/v1/ph/pending-staff/{profile_id}/activate",
        json={"email": "valerie.standingready@crbcl.ca", "position": "Prevention Worker"},
        headers=headers,
    )
    assert act_res.status_code == 200
    activated_data = act_res.json()
    assert activated_data["onboarding_status"] == "ACTIVATED"
    assert activated_data["employee_id"] is not None

    # 5. Verify Employee record was created in employees table
    emp_res = await db_session.execute(select(Employee).where(Employee.email == "valerie.standingready@crbcl.ca"))
    emp = emp_res.scalars().first()
    assert emp is not None
    assert emp.first_name == "Valerie"
    assert emp.last_name == "Standingready"
    assert emp.department == "Prevention & Healing"

    # 6. Verify historical EnhancementProgram now links to new Employee
    await db_session.refresh(ep)
    assert ep.employee_id == emp.id


@pytest.mark.asyncio
async def test_client_reconciliation_linking(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    """Verifies:

    - Legacy RedMane ID is registered in reconciliation queue
    - Confirming matches to canonical 10-digit Person ID
    - Rejecting updates status to REJECTED
    """
    headers = executive_director_user["headers"]

    # Create a canonical Person with 10-digit person_id_number
    person = Person(
        person_id_number="1234567890",
        first_name="George",
        last_name="Pelletier",
    )
    db_session.add(person)
    await db_session.commit()
    await db_session.refresh(person)

    # Register legacy client ref
    cref = PHClientRef(
        legacy_id="RM-88219",
        legacy_id_type="REDMANE",
        status="UNRESOLVED",
    )
    db_session.add(cref)
    await db_session.commit()
    await db_session.refresh(cref)
    cref_id = str(cref.id)

    # 1. Search canonical persons
    search_res = await client.get("/api/v1/ph/persons/search?q=George", headers=headers)
    assert search_res.status_code == 200
    search_data = search_res.json()
    assert len(search_data) >= 1
    assert any(p["person_id_number"] == "1234567890" for p in search_data)

    # Also search by 10-digit number
    num_search = await client.get("/api/v1/ph/persons/search?q=1234567890", headers=headers)
    assert num_search.status_code == 200
    assert len(num_search.json()) == 1

    # Create a FamilySupportRecord referencing legacy ID RM-88219
    fsr = FamilySupportRecord(
        program="Family Support",
        client_ref_legacy_id="RM-88219",
        client_match_status="UNRESOLVED",
        fiscal_year="2026-2027",
    )
    db_session.add(fsr)
    await db_session.commit()
    await db_session.refresh(fsr)
    assert fsr.person_id is None
    assert fsr.client_match_status == "UNRESOLVED"

    # 2. List unresolved client refs
    list_res = await client.get("/api/v1/ph/client-refs?status=UNRESOLVED", headers=headers)
    assert list_res.status_code == 200
    assert any(r["id"] == cref_id for r in list_res.json())

    # 3. Confirm link to canonical Person via reconcile endpoint
    confirm_res = await client.post(
        f"/api/v1/ph/client-refs/{cref_id}/reconcile",
        json={"action": "CONFIRM", "person_id": str(person.id)},
        headers=headers,
    )
    assert confirm_res.status_code == 200
    assert confirm_res.json()["status"] == "CONFIRMED"
    assert confirm_res.json()["person_id"] == str(person.id)
    assert confirm_res.json()["person_id_number"] == "1234567890"

    # Verify FamilySupportRecord was updated with the canonical Person link
    await db_session.refresh(fsr)
    assert fsr.person_id == person.id
    assert fsr.client_match_status == "CONFIRMED"

    # 4. Reject a different client ref via reconcile endpoint
    cref2 = PHClientRef(legacy_id="RM-INVALID", legacy_id_type="REDMANE", status="UNRESOLVED")
    db_session.add(cref2)
    await db_session.commit()
    await db_session.refresh(cref2)

    rej_res = await client.post(
        f"/api/v1/ph/client-refs/{cref2.id}/reconcile",
        json={"action": "REJECT"},
        headers=headers,
    )
    assert rej_res.status_code == 200
    assert rej_res.json()["status"] == "REJECTED"


@pytest.mark.asyncio
async def test_excel_import_wizard_analyze_approve_execute(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    """End-to-end test of the 8-step Excel Import Wizard using synthetic data:

    1. Generate synthetic workbook in memory with Enhancement Programs headers
    2. Upload to /ph/import/analyze (Step 2 & 3: Validate & preview)
    3. Verify batch in AWAITING_APPROVAL (Step 6 & 7)
    4. Approve batch /ph/import/approve/{batch_id}
    5. Execute batch /ph/import/execute/{batch_id} (Step 8: Transact)
    6. Verify records imported into database
    """
    headers = executive_director_user["headers"]

    # 1. Create synthetic workbook in memory
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enhancement Programs"

    # Set up row 7 headers matching real tracker
    headers_row = [
        "", "Date\n(dd-Mon-yy)", "Employee Name", "Employee Position", "Program", "Program Activity",
        "Date Provided\n(dd-Mon-yy)", "City\n(e.g. Regina)", "Province\n(e.g. SK)", "Partners/Collaborators",
        "Closing Date", "In Province Families", "Out of Province Families", "Total Families", "Prev",
        "Children \n0-12", "Youth\n13-17", "Post Majority\n18-26", "Parents/Caregivers", "Family Members",
        "Others", "Total Participants", "Account", "Sub", "Gift Card", "Purchase Order", "EFT", "Cheque",
        "Credit Card", "Courtesy Card", "Total Cost", "Comments"
    ]
    for c_idx, h in enumerate(headers_row, 1):
        ws.cell(7, c_idx, h)

    # Add 2 synthetic rows (row 8 & row 9)
    # Row 8: Valid row
    ws.cell(8, 2, "2026-05-10")
    ws.cell(8, 3, "Synthetic Staff One")
    ws.cell(8, 4, "Facilitator")
    ws.cell(8, 5, "Youth Healing Circle")
    ws.cell(8, 6, "Cultural Workshop")
    ws.cell(8, 8, "Regina")
    ws.cell(8, 9, "SK")
    ws.cell(8, 12, 5)   # in province fam
    ws.cell(8, 13, 1)   # out of province fam
    ws.cell(8, 16, 6)   # children
    ws.cell(8, 17, 4)   # youth
    ws.cell(8, 23, "41000")
    ws.cell(8, 24, "213")
    ws.cell(8, 25, 100.00) # gift card
    ws.cell(8, 26, 200.00) # PO
    ws.cell(8, 31, 300.00) # total cost
    ws.cell(8, 32, "Synthetic test row 1")

    # Row 9: Valid row
    ws.cell(9, 2, "2026-05-12")
    ws.cell(9, 3, "Synthetic Staff Two")
    ws.cell(9, 4, "Counsellor")
    ws.cell(9, 5, "Parenting Wisdom")
    ws.cell(9, 8, "Saskatoon")
    ws.cell(9, 9, "SK")
    ws.cell(9, 12, 8)
    ws.cell(9, 16, 10)
    ws.cell(9, 23, "41000")
    ws.cell(9, 24, "215")
    ws.cell(9, 27, 450.00) # EFT
    ws.cell(9, 31, 450.00)
    ws.cell(9, 32, "Synthetic test row 2")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # 2. Upload to analyze endpoint
    files = {"file": ("synthetic_tracker.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    analyze_res = await client.post("/api/v1/ph/import/analyze", files=files, headers=headers)
    assert analyze_res.status_code == 200, analyze_res.text
    analyze_data = analyze_res.json()

    batch_id = analyze_data["batch_id"]
    assert analyze_data["total_rows_analyzed"] == 2
    assert analyze_data["valid_rows"] == 2
    assert analyze_data["invalid_rows"] == 0
    assert analyze_data["ready_for_approval"] is True
    assert "Synthetic Staff One" in analyze_data["employee_matching_issues"]
    assert "Synthetic Staff Two" in analyze_data["employee_matching_issues"]

    # 3. Verify cannot execute before approval
    exec_fail = await client.post(f"/api/v1/ph/import/execute/{batch_id}", headers=headers)
    assert exec_fail.status_code == 400
    assert "APPROVED" in exec_fail.text

    # 4. Approve batch
    appr_res = await client.post(f"/api/v1/ph/import/approve/{batch_id}", headers=headers)
    assert appr_res.status_code == 200
    assert appr_res.json()["status"] == "APPROVED"

    # 5. Execute import
    exec_res = await client.post(f"/api/v1/ph/import/execute/{batch_id}", headers=headers)
    assert exec_res.status_code == 200
    exec_data = exec_res.json()
    assert exec_data["status"] == "COMPLETED"
    assert exec_data["imported_rows"] == 2

    # 6. Verify records exist in EnhancementProgram table
    db_records = await db_session.execute(
        select(EnhancementProgram).where(EnhancementProgram.import_batch_id == uuid.UUID(batch_id))
    )
    records = db_records.scalars().all()
    assert len(records) == 2
    assert {r.program for r in records} == {"Youth Healing Circle", "Parenting Wisdom"}

    # 7. Test cancelling a newly analyzed batch
    analyze_res2 = await client.post("/api/v1/ph/import/analyze", files=files, headers=headers)
    assert analyze_res2.status_code == 200
    batch_id2 = analyze_res2.json()["batch_id"]

    cancel_res = await client.post(f"/api/v1/ph/import/cancel/{batch_id2}", headers=headers)
    assert cancel_res.status_code == 200
    assert cancel_res.json()["status"] == "CANCELLED"


@pytest.mark.asyncio
async def test_annual_summary_reporting(
    client: AsyncClient, executive_director_user, db_session: AsyncSession
):
    """Verifies that the /ph/summary endpoint aggregates totals accurately across reporting streams."""
    headers = executive_director_user["headers"]

    # Seed one of each record
    ep = EnhancementProgram(
        program="Annual Culture Gathering",
        fiscal_year="2026-2027",
        families_in_province=10,
        families_out_of_province=2,
        participants_children_0_12=20,
        participants_youth_13_17=10,
        expense_gift_card=Decimal("500.00"),
        expense_purchase_order=Decimal("1000.00"),
    )
    fsr = FamilySupportRecord(
        program="Direct Family Support",
        fiscal_year="2026-2027",
        expense_gift_card=Decimal("300.00"),
        expense_cheque=Decimal("200.00"),
    )
    ons = OtherNationsSupport(
        service_type="Eagle Nest Lodging",
        fiscal_year="2026-2027",
        cost=Decimal("750.00"),
    )
    db_session.add_all([ep, fsr, ons])
    await db_session.commit()

    res = await client.get("/api/v1/ph/summary?fiscal_year=2026-2027", headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert data["fiscal_year"] == "2026-2027"
    assert data["enhancement_programs"]["total_cost"] == 1500.00
    assert data["enhancement_programs"]["total_families"] == 12
    assert data["family_support"]["total_cost"] == 500.00
    assert data["other_nations_support"]["total_cost"] == 750.00
    assert data["combined_total_cost"] == 2750.00


@pytest.mark.asyncio
async def test_security_it_admin_boundary(
    client: AsyncClient, it_admin_user
):
    """IT Admin must NOT have clinical or Prevention & Healing record management access.

    Verifies the principle of least privilege & privacy boundaries.
    """
    headers = it_admin_user["headers"]

    # IT Admin attempting to read PH records
    res = await client.get("/api/v1/ph/enhancement-programs", headers=headers)
    assert res.status_code == 403

    # IT Admin attempting to execute import
    res2 = await client.post(f"/api/v1/ph/import/execute/{uuid.uuid4()}", headers=headers)
    assert res2.status_code == 403
