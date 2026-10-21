"""Prevention & Healing Records Workspace — API routes.

Endpoints:
  Enhancement Programs
    GET    /ph/enhancement-programs
    POST   /ph/enhancement-programs
    GET    /ph/enhancement-programs/{id}
    PUT    /ph/enhancement-programs/{id}
    DELETE /ph/enhancement-programs/{id}

  Family Support Records
    GET    /ph/family-support
    POST   /ph/family-support
    GET    /ph/family-support/{id}
    PUT    /ph/family-support/{id}
    DELETE /ph/family-support/{id}

  Other Nations Support
    GET    /ph/other-nations
    POST   /ph/other-nations
    GET    /ph/other-nations/{id}
    PUT    /ph/other-nations/{id}

  Pending Staff Onboarding
    GET    /ph/pending-staff
    GET    /ph/pending-staff/{id}
    POST   /ph/pending-staff
    PUT    /ph/pending-staff/{id}

  Client Reconciliation
    GET    /ph/client-refs
    PUT    /ph/client-refs/{id}/confirm
    PUT    /ph/client-refs/{id}/reject

  Import (Excel wizard)
    POST   /ph/import/analyze       — analyze without importing
    POST   /ph/import/approve/{batch_id}
    POST   /ph/import/execute/{batch_id}
    GET    /ph/import/batches
    GET    /ph/import/batches/{batch_id}
    GET    /ph/import/batches/{batch_id}/rows

  Reports / Summary
    GET    /ph/summary               — annual summary & stats
"""

from __future__ import annotations

import uuid
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from io import BytesIO
from pydantic import BaseModel
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl

from app.core.database import get_db
from app.models.org_ops import Employee
from app.models.person import Person
from app.models.prevention_healing import (
    EnhancementProgram,
    FamilySupportRecord,
    OtherNationsSupport,
    PHClientRef,
    PHImportBatch,
    PHImportRow,
    PendingStaffProfile,
)
from app.models.user import User
from app.permissions.constants import Permissions
from app.permissions.dependencies import require_permission

router = APIRouter(prefix="/ph", tags=["Prevention & Healing"])


# ──────────────────────────────────────────────────────────────
# Shared helpers
# ──────────────────────────────────────────────────────────────

def _not_found(entity: str, id_: uuid.UUID) -> HTTPException:
    return HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"{entity} {id_} not found")


def _decimal(v) -> Decimal:
    return Decimal(str(v)) if v is not None else Decimal("0.00")


# ──────────────────────────────────────────────────────────────
# Pydantic schemas
# ──────────────────────────────────────────────────────────────

class EnhancementProgramIn(BaseModel):
    employee_id: uuid.UUID | None = None
    pending_staff_id: uuid.UUID | None = None
    employee_name_raw: str | None = None
    position_raw: str | None = None
    record_date: date | None = None
    program: str
    program_activity: str | None = None
    date_provided: date | None = None
    closing_date: date | None = None
    city: str | None = None
    province: str | None = None
    partners: str | None = None
    families_in_province: int = 0
    families_out_of_province: int = 0
    participants_children_0_12: int = 0
    participants_youth_13_17: int = 0
    participants_post_majority_18_26: int = 0
    participants_parents_caregivers: int = 0
    participants_family_members: int = 0
    participants_others: int = 0
    account_code: str | None = None
    subcode: str | None = None
    expense_gift_card: Decimal = Decimal("0.00")
    expense_purchase_order: Decimal = Decimal("0.00")
    expense_eft: Decimal = Decimal("0.00")
    expense_cheque: Decimal = Decimal("0.00")
    expense_credit_card: Decimal = Decimal("0.00")
    expense_courtesy_card: Decimal = Decimal("0.00")
    comments: str | None = None
    fiscal_year: str = "2026-2027"


class FamilySupportIn(BaseModel):
    employee_id: uuid.UUID | None = None
    pending_staff_id: uuid.UUID | None = None
    employee_name_raw: str | None = None
    position_raw: str | None = None
    service_area: str | None = None
    record_date: date | None = None
    program: str
    support_type: str | None = None
    description: str | None = None
    partners: str | None = None
    date_provided: date | None = None
    person_id: uuid.UUID | None = None
    client_ref_legacy_id: str | None = None
    city: str | None = None
    province: str | None = None
    participants_children_0_12: int = 0
    participants_youth_13_17: int = 0
    participants_post_majority_18_26: int = 0
    participants_adults_18_plus: int = 0
    account_code: str | None = None
    subcode: str | None = None
    project_code: str | None = None
    expense_gift_card: Decimal = Decimal("0.00")
    expense_purchase_order: Decimal = Decimal("0.00")
    expense_eft: Decimal = Decimal("0.00")
    expense_cheque: Decimal = Decimal("0.00")
    expense_credit_card: Decimal = Decimal("0.00")
    expense_petty_cash: Decimal = Decimal("0.00")
    expense_courtesy_card: Decimal = Decimal("0.00")
    comments: str | None = None
    fiscal_year: str = "2026-2027"
    client_match_status: str = "UNRESOLVED"


class OtherNationsIn(BaseModel):
    employee_id: uuid.UUID | None = None
    pending_staff_id: uuid.UUID | None = None
    worker_name_raw: str | None = None
    record_date: date | None = None
    family_reference_raw: str | None = None
    service_type: str | None = None
    cost: Decimal = Decimal("0.00")
    on_reserve: bool | None = None
    cowessess_connection: str | None = None
    other_community: str | None = None
    historical_reference_number: str | None = None
    comments: str | None = None
    person_id: uuid.UUID | None = None
    fiscal_year: str = "2026-2027"


class PendingStaffIn(BaseModel):
    first_name: str
    last_name: str
    display_name_from_excel: str
    position: str | None = None
    department: str | None = None
    email: str | None = None
    phone: str | None = None
    notes: str | None = None
    onboarding_status: str = "PENDING"


class PendingStaffUpdate(BaseModel):
    email: str | None = None
    phone: str | None = None
    position: str | None = None
    department: str | None = None
    onboarding_status: str | None = None
    employee_id: uuid.UUID | None = None
    notes: str | None = None


# ──────────────────────────────────────────────────────────────
# Enhancement Programs
# ──────────────────────────────────────────────────────────────

@router.get("/enhancement-programs")
async def list_enhancement_programs(
    fiscal_year: str | None = Query(None),
    program: str | None = Query(None),
    employee_id: uuid.UUID | None = Query(None),
    limit: int = Query(50, le=200),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_READ)),
):
    q = select(EnhancementProgram)
    if fiscal_year:
        q = q.where(EnhancementProgram.fiscal_year == fiscal_year)
    if program:
        q = q.where(EnhancementProgram.program.ilike(f"%{program}%"))
    if employee_id:
        q = q.where(EnhancementProgram.employee_id == employee_id)
    q = q.order_by(EnhancementProgram.record_date.desc().nullslast()).offset(offset).limit(limit)
    result = await db.execute(q)
    rows = result.scalars().all()
    return [_serialize_ep(r) for r in rows]


@router.get("/enhancement-programs/{record_id}")
async def get_enhancement_program(
    record_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_READ)),
):
    row = await db.get(EnhancementProgram, record_id)
    if not row:
        raise _not_found("Enhancement program", record_id)
    return _serialize_ep(row)


@router.post("/enhancement-programs", status_code=201)
async def create_enhancement_program(
    body: EnhancementProgramIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_CREATE)),
):
    rec = EnhancementProgram(**body.model_dump())
    db.add(rec)
    await db.commit()
    await db.refresh(rec)
    return _serialize_ep(rec)


@router.put("/enhancement-programs/{record_id}")
async def update_enhancement_program(
    record_id: uuid.UUID,
    body: EnhancementProgramIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_UPDATE)),
):
    row = await db.get(EnhancementProgram, record_id)
    if not row:
        raise _not_found("Enhancement program", record_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(row, k, v)
    await db.commit()
    await db.refresh(row)
    return _serialize_ep(row)


@router.delete("/enhancement-programs/{record_id}", status_code=204)
async def delete_enhancement_program(
    record_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_DELETE)),
):
    row = await db.get(EnhancementProgram, record_id)
    if not row:
        raise _not_found("Enhancement program", record_id)
    await db.delete(row)
    await db.commit()


def _serialize_ep(r: EnhancementProgram) -> dict:
    return {
        "id": str(r.id),
        "employee_id": str(r.employee_id) if r.employee_id else None,
        "pending_staff_id": str(r.pending_staff_id) if r.pending_staff_id else None,
        "employee_name_raw": r.employee_name_raw,
        "position_raw": r.position_raw,
        "record_date": r.record_date.isoformat() if r.record_date else None,
        "program": r.program,
        "program_activity": r.program_activity,
        "date_provided": r.date_provided.isoformat() if r.date_provided else None,
        "closing_date": r.closing_date.isoformat() if r.closing_date else None,
        "city": r.city,
        "province": r.province,
        "partners": r.partners,
        "families_in_province": r.families_in_province,
        "families_out_of_province": r.families_out_of_province,
        "total_families": r.total_families,
        "participants_children_0_12": r.participants_children_0_12,
        "participants_youth_13_17": r.participants_youth_13_17,
        "participants_post_majority_18_26": r.participants_post_majority_18_26,
        "participants_parents_caregivers": r.participants_parents_caregivers,
        "participants_family_members": r.participants_family_members,
        "participants_others": r.participants_others,
        "total_participants": r.total_participants,
        "account_code": r.account_code,
        "subcode": r.subcode,
        "expense_gift_card": float(r.expense_gift_card),
        "expense_purchase_order": float(r.expense_purchase_order),
        "expense_eft": float(r.expense_eft),
        "expense_cheque": float(r.expense_cheque),
        "expense_credit_card": float(r.expense_credit_card),
        "expense_courtesy_card": float(r.expense_courtesy_card),
        "total_cost": float(r.total_cost),
        "comments": r.comments,
        "fiscal_year": r.fiscal_year,
        "import_batch_id": str(r.import_batch_id) if r.import_batch_id else None,
        "source_row_number": r.source_row_number,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


# ──────────────────────────────────────────────────────────────
# Family Support Records
# ──────────────────────────────────────────────────────────────

@router.get("/family-support")
async def list_family_support(
    fiscal_year: str | None = Query(None),
    program: str | None = Query(None),
    support_type: str | None = Query(None),
    service_area: str | None = Query(None),
    employee_id: uuid.UUID | None = Query(None),
    client_match_status: str | None = Query(None),
    limit: int = Query(50, le=200),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_READ)),
):
    q = select(FamilySupportRecord)
    if fiscal_year:
        q = q.where(FamilySupportRecord.fiscal_year == fiscal_year)
    if program:
        q = q.where(FamilySupportRecord.program.ilike(f"%{program}%"))
    if support_type:
        q = q.where(FamilySupportRecord.support_type.ilike(f"%{support_type}%"))
    if service_area:
        q = q.where(FamilySupportRecord.service_area == service_area)
    if employee_id:
        q = q.where(FamilySupportRecord.employee_id == employee_id)
    if client_match_status:
        q = q.where(FamilySupportRecord.client_match_status == client_match_status)
    q = q.order_by(FamilySupportRecord.record_date.desc().nullslast()).offset(offset).limit(limit)
    result = await db.execute(q)
    return [_serialize_fsr(r) for r in result.scalars().all()]


@router.get("/family-support/{record_id}")
async def get_family_support(
    record_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_READ)),
):
    row = await db.get(FamilySupportRecord, record_id)
    if not row:
        raise _not_found("Family support record", record_id)
    return _serialize_fsr(row)


@router.post("/family-support", status_code=201)
async def create_family_support(
    body: FamilySupportIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_CREATE)),
):
    rec = FamilySupportRecord(**body.model_dump())
    db.add(rec)
    await db.commit()
    await db.refresh(rec)
    return _serialize_fsr(rec)


@router.put("/family-support/{record_id}")
async def update_family_support(
    record_id: uuid.UUID,
    body: FamilySupportIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_UPDATE)),
):
    row = await db.get(FamilySupportRecord, record_id)
    if not row:
        raise _not_found("Family support record", record_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(row, k, v)
    await db.commit()
    await db.refresh(row)
    return _serialize_fsr(row)


@router.delete("/family-support/{record_id}", status_code=204)
async def delete_family_support(
    record_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_DELETE)),
):
    row = await db.get(FamilySupportRecord, record_id)
    if not row:
        raise _not_found("Family support record", record_id)
    await db.delete(row)
    await db.commit()


def _serialize_fsr(r: FamilySupportRecord) -> dict:
    return {
        "id": str(r.id),
        "employee_id": str(r.employee_id) if r.employee_id else None,
        "pending_staff_id": str(r.pending_staff_id) if r.pending_staff_id else None,
        "employee_name_raw": r.employee_name_raw,
        "position_raw": r.position_raw,
        "service_area": r.service_area,
        "record_date": r.record_date.isoformat() if r.record_date else None,
        "program": r.program,
        "support_type": r.support_type,
        "description": r.description,
        "partners": r.partners,
        "date_provided": r.date_provided.isoformat() if r.date_provided else None,
        "person_id": str(r.person_id) if r.person_id else None,
        "client_ref_legacy_id": r.client_ref_legacy_id,
        "city": r.city,
        "province": r.province,
        "participants_children_0_12": r.participants_children_0_12,
        "participants_youth_13_17": r.participants_youth_13_17,
        "participants_post_majority_18_26": r.participants_post_majority_18_26,
        "participants_adults_18_plus": r.participants_adults_18_plus,
        "account_code": r.account_code,
        "subcode": r.subcode,
        "project_code": r.project_code,
        "expense_gift_card": float(r.expense_gift_card),
        "expense_purchase_order": float(r.expense_purchase_order),
        "expense_eft": float(r.expense_eft),
        "expense_cheque": float(r.expense_cheque),
        "expense_credit_card": float(r.expense_credit_card),
        "expense_petty_cash": float(r.expense_petty_cash),
        "expense_courtesy_card": float(r.expense_courtesy_card),
        "total_cost": float(r.total_cost),
        "comments": r.comments,
        "fiscal_year": r.fiscal_year,
        "client_match_status": r.client_match_status,
        "import_batch_id": str(r.import_batch_id) if r.import_batch_id else None,
        "source_row_number": r.source_row_number,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


# ──────────────────────────────────────────────────────────────
# Other Nations Support
# ──────────────────────────────────────────────────────────────

@router.get("/other-nations")
async def list_other_nations(
    fiscal_year: str | None = Query(None),
    service_type: str | None = Query(None),
    limit: int = Query(50, le=200),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_READ)),
):
    q = select(OtherNationsSupport)
    if fiscal_year:
        q = q.where(OtherNationsSupport.fiscal_year == fiscal_year)
    if service_type:
        q = q.where(OtherNationsSupport.service_type.ilike(f"%{service_type}%"))
    q = q.order_by(OtherNationsSupport.record_date.desc().nullslast()).offset(offset).limit(limit)
    result = await db.execute(q)
    return [_serialize_ons(r) for r in result.scalars().all()]


@router.get("/other-nations/{record_id}")
async def get_other_nations(
    record_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_READ)),
):
    row = await db.get(OtherNationsSupport, record_id)
    if not row:
        raise _not_found("Other Nations support record", record_id)
    return _serialize_ons(row)


@router.post("/other-nations", status_code=201)
async def create_other_nations(
    body: OtherNationsIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_CREATE)),
):
    rec = OtherNationsSupport(**body.model_dump())
    db.add(rec)
    await db.commit()
    await db.refresh(rec)
    return _serialize_ons(rec)


@router.put("/other-nations/{record_id}")
async def update_other_nations(
    record_id: uuid.UUID,
    body: OtherNationsIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_RECORD_UPDATE)),
):
    row = await db.get(OtherNationsSupport, record_id)
    if not row:
        raise _not_found("Other Nations record", record_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(row, k, v)
    await db.commit()
    await db.refresh(row)
    return _serialize_ons(row)


def _serialize_ons(r: OtherNationsSupport) -> dict:
    return {
        "id": str(r.id),
        "employee_id": str(r.employee_id) if r.employee_id else None,
        "pending_staff_id": str(r.pending_staff_id) if r.pending_staff_id else None,
        "worker_name_raw": r.worker_name_raw,
        "record_date": r.record_date.isoformat() if r.record_date else None,
        "family_reference_raw": r.family_reference_raw,
        "service_type": r.service_type,
        "cost": float(r.cost),
        "on_reserve": r.on_reserve,
        "cowessess_connection": r.cowessess_connection,
        "other_community": r.other_community,
        "historical_reference_number": r.historical_reference_number,
        "comments": r.comments,
        "person_id": str(r.person_id) if r.person_id else None,
        "client_match_status": r.client_match_status,
        "fiscal_year": r.fiscal_year,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


# ──────────────────────────────────────────────────────────────
# Pending Staff Onboarding
# ──────────────────────────────────────────────────────────────

@router.get("/pending-staff")
async def list_pending_staff(
    status_filter: str | None = Query(None, alias="status"),
    limit: int = Query(50, le=200),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_STAFF_ONBOARD_READ)),
):
    q = select(PendingStaffProfile)
    if status_filter:
        q = q.where(PendingStaffProfile.onboarding_status == status_filter)
    q = q.order_by(PendingStaffProfile.last_name).offset(offset).limit(limit)
    result = await db.execute(q)
    return [_serialize_ps(r) for r in result.scalars().all()]


@router.get("/pending-staff/{profile_id}")
async def get_pending_staff(
    profile_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_STAFF_ONBOARD_READ)),
):
    row = await db.get(PendingStaffProfile, profile_id)
    if not row:
        raise _not_found("Pending staff profile", profile_id)
    return _serialize_ps(row)


@router.post("/pending-staff", status_code=201)
async def create_pending_staff(
    body: PendingStaffIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_STAFF_ONBOARD_MANAGE)),
):
    rec = PendingStaffProfile(**body.model_dump())
    db.add(rec)
    await db.commit()
    await db.refresh(rec)
    return _serialize_ps(rec)


@router.put("/pending-staff/{profile_id}")
async def update_pending_staff(
    profile_id: uuid.UUID,
    body: PendingStaffUpdate,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_STAFF_ONBOARD_MANAGE)),
):
    row = await db.get(PendingStaffProfile, profile_id)
    if not row:
        raise _not_found("Pending staff profile", profile_id)
    for k, v in body.model_dump(exclude_unset=True, exclude_none=False).items():
        if v is not None:
            setattr(row, k, v)
    await db.commit()
    await db.refresh(row)
    return _serialize_ps(row)


class PendingStaffActivateIn(BaseModel):
    email: str | None = None
    position: str | None = None
    department: str = "Prevention & Healing"


@router.post("/pending-staff/{profile_id}/activate")
async def activate_pending_staff(
    profile_id: uuid.UUID,
    body: PendingStaffActivateIn | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission(Permissions.PH_STAFF_ONBOARD_MANAGE)),
):
    row = await db.get(PendingStaffProfile, profile_id)
    if not row:
        raise _not_found("Pending staff profile", profile_id)

    email = (body.email if body and body.email else row.email)
    if not email or not email.strip() or "@" not in email:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="A verified email address is strictly required to activate an employee profile. Inactive profiles remain pending until verified."
        )
    email = email.strip().lower()

    # Search existing Employee directory first (by email or full name)
    q_emp = select(Employee).where(
        (func.lower(Employee.email) == email) |
        (
            (func.lower(Employee.first_name) == row.first_name.lower()) &
            (func.lower(Employee.last_name) == row.last_name.lower())
        )
    )
    res_emp = await db.execute(q_emp)
    existing_emp = res_emp.scalars().first()

    if not existing_emp:
        pos = (body.position if body and body.position else row.position) or "Prevention Worker"
        dept = (body.department if body and body.department else row.department) or "Prevention & Healing"
        emp_num = f"PH-{uuid.uuid4().hex[:6].upper()}"
        existing_emp = Employee(
            first_name=row.first_name,
            last_name=row.last_name,
            email=email,
            phone=row.phone,
            position=pos,
            department=dept,
            employee_number=emp_num,
            hire_date=date.today(),
            employment_status="ACTIVE",
        )
        db.add(existing_emp)
        await db.flush()
    else:
        existing_emp.email = email
        if row.phone and not existing_emp.phone:
            existing_emp.phone = row.phone

    row.email = email
    row.employee_id = existing_emp.id
    row.onboarding_status = "ACTIVATED"

    # Link historical records associated with this pending staff profile
    await db.execute(
        update(EnhancementProgram)
        .where(EnhancementProgram.pending_staff_id == row.id)
        .values(employee_id=existing_emp.id)
    )
    await db.execute(
        update(FamilySupportRecord)
        .where(FamilySupportRecord.pending_staff_id == row.id)
        .values(employee_id=existing_emp.id)
    )
    await db.execute(
        update(OtherNationsSupport)
        .where(OtherNationsSupport.pending_staff_id == row.id)
        .values(employee_id=existing_emp.id)
    )

    await db.commit()
    await db.refresh(row)
    return _serialize_ps(row)


def _serialize_ps(r: PendingStaffProfile) -> dict:
    return {
        "id": str(r.id),
        "first_name": r.first_name,
        "last_name": r.last_name,
        "display_name_from_excel": r.display_name_from_excel,
        "position": r.position,
        "department": r.department,
        "email": r.email,
        "phone": r.phone,
        "onboarding_status": r.onboarding_status,
        "employee_id": str(r.employee_id) if r.employee_id else None,
        "notes": r.notes,
        "source_batch_id": str(r.source_batch_id) if r.source_batch_id else None,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


# ──────────────────────────────────────────────────────────────
# Client Reconciliation
# ──────────────────────────────────────────────────────────────

class ClientRefConfirmIn(BaseModel):
    person_id: uuid.UUID


class ClientRefReconcileIn(BaseModel):
    action: str  # "CONFIRM" | "REJECT"
    person_id: uuid.UUID | None = None


@router.get("/persons/search")
async def search_persons(
    q: str = Query(..., min_length=1),
    limit: int = Query(20, le=50),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_CLIENT_RECONCILE)),
):
    """Search canonical Person database for matching client identities."""
    term = f"%{q.strip()}%"
    stmt = (
        select(Person)
        .where(
            (Person.person_id_number.ilike(term))
            | (Person.first_name.ilike(term))
            | (Person.last_name.ilike(term))
            | (func.concat(Person.first_name, " ", Person.last_name).ilike(term))
        )
        .order_by(Person.last_name, Person.first_name)
        .limit(limit)
    )
    result = await db.execute(stmt)
    persons = result.scalars().all()
    return [
        {
            "id": str(p.id),
            "person_id_number": p.person_id_number,
            "first_name": p.first_name,
            "last_name": p.last_name,
            "preferred_name": p.preferred_name,
            "date_of_birth": p.date_of_birth.isoformat() if p.date_of_birth else None,
            "band_nation": p.band_nation,
        }
        for p in persons
    ]


@router.get("/client-refs")
async def list_client_refs(
    status_filter: str | None = Query(None, alias="status"),
    limit: int = Query(50, le=200),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_CLIENT_RECONCILE)),
):
    q = select(PHClientRef)
    if status_filter:
        q = q.where(PHClientRef.status == status_filter)
    q = q.order_by(PHClientRef.created_at.desc()).offset(offset).limit(limit)
    result = await db.execute(q)
    return [_serialize_cr(r) for r in result.scalars().all()]


@router.put("/client-refs/{ref_id}/confirm")
async def confirm_client_ref(
    ref_id: uuid.UUID,
    body: ClientRefConfirmIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_CLIENT_RECONCILE)),
):
    row = await db.get(PHClientRef, ref_id)
    if not row:
        raise _not_found("Client reference", ref_id)
    person = await db.get(Person, body.person_id)
    if not person:
        raise HTTPException(status_code=404, detail="Canonical Person not found")
    row.person_id = body.person_id
    row.status = "CONFIRMED"
    row.confirmed_at = datetime.utcnow()
    # Cascade to FamilySupportRecord
    await db.execute(
        update(FamilySupportRecord)
        .where(FamilySupportRecord.client_ref_legacy_id == row.legacy_id)
        .values(person_id=body.person_id, client_match_status="CONFIRMED")
    )
    await db.commit()
    await db.refresh(row)
    return _serialize_cr(row)


@router.put("/client-refs/{ref_id}/reject")
async def reject_client_ref(
    ref_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_CLIENT_RECONCILE)),
):
    row = await db.get(PHClientRef, ref_id)
    if not row:
        raise _not_found("Client reference", ref_id)
    row.status = "REJECTED"
    await db.execute(
        update(FamilySupportRecord)
        .where(FamilySupportRecord.client_ref_legacy_id == row.legacy_id)
        .values(client_match_status="REJECTED")
    )
    await db.commit()
    await db.refresh(row)
    return _serialize_cr(row)


@router.post("/client-refs/{ref_id}/reconcile")
async def reconcile_client_ref(
    ref_id: uuid.UUID,
    body: ClientRefReconcileIn,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_CLIENT_RECONCILE)),
):
    row = await db.get(PHClientRef, ref_id)
    if not row:
        raise _not_found("Client reference", ref_id)

    action = (body.action or "").strip().upper()
    if action == "CONFIRM":
        if not body.person_id:
            raise HTTPException(status_code=400, detail="person_id is required to confirm a client reference.")
        person = await db.get(Person, body.person_id)
        if not person:
            raise HTTPException(status_code=404, detail=f"Canonical Person {body.person_id} not found.")
        row.person_id = body.person_id
        row.status = "CONFIRMED"
        row.confirmed_at = datetime.utcnow()
        await db.execute(
            update(FamilySupportRecord)
            .where(FamilySupportRecord.client_ref_legacy_id == row.legacy_id)
            .values(person_id=body.person_id, client_match_status="CONFIRMED")
        )
    elif action == "REJECT":
        row.status = "REJECTED"
        await db.execute(
            update(FamilySupportRecord)
            .where(FamilySupportRecord.client_ref_legacy_id == row.legacy_id)
            .values(client_match_status="REJECTED")
        )
    else:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid reconciliation action {body.action!r}. Must be 'CONFIRM' or 'REJECT'."
        )

    await db.commit()
    await db.refresh(row)
    return _serialize_cr(row)


def _serialize_cr(r: PHClientRef) -> dict:
    person_obj = getattr(r, "person", None)
    return {
        "id": str(r.id),
        "legacy_id": r.legacy_id,
        "legacy_id_type": r.legacy_id_type,
        "person_id": str(r.person_id) if r.person_id else None,
        "person_id_number": person_obj.person_id_number if person_obj else None,
        "person_name": f"{person_obj.first_name} {person_obj.last_name}" if person_obj else None,
        "status": r.status,
        "confirmed_at": r.confirmed_at.isoformat() if r.confirmed_at else None,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


# ──────────────────────────────────────────────────────────────
# Import Batches
# ──────────────────────────────────────────────────────────────

@router.get("/import/batches")
async def list_import_batches(
    status_filter: str | None = Query(None, alias="status"),
    limit: int = Query(20, le=100),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_IMPORT_HISTORY_READ)),
):
    q = select(PHImportBatch)
    if status_filter:
        q = q.where(PHImportBatch.status == status_filter)
    q = q.order_by(PHImportBatch.created_at.desc()).offset(offset).limit(limit)
    result = await db.execute(q)
    return [_serialize_batch(r) for r in result.scalars().all()]


@router.get("/import/batches/{batch_id}")
async def get_import_batch(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_IMPORT_HISTORY_READ)),
):
    row = await db.get(PHImportBatch, batch_id)
    if not row:
        raise _not_found("Import batch", batch_id)
    return _serialize_batch(row)


@router.get("/import/batches/{batch_id}/rows")
async def get_import_batch_rows(
    batch_id: uuid.UUID,
    status_filter: str | None = Query(None, alias="status"),
    limit: int = Query(100, le=500),
    offset: int = Query(0),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_IMPORT_HISTORY_READ)),
):
    q = select(PHImportRow).where(PHImportRow.batch_id == batch_id)
    if status_filter:
        q = q.where(PHImportRow.status == status_filter)
    q = q.order_by(PHImportRow.source_row_number).offset(offset).limit(limit)
    result = await db.execute(q)
    return [_serialize_row(r) for r in result.scalars().all()]


@router.post("/import/approve/{batch_id}")
async def approve_import(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_IMPORT_APPROVE)),
):
    batch = await db.get(PHImportBatch, batch_id)
    if not batch:
        raise _not_found("Import batch", batch_id)
    if batch.status != "AWAITING_APPROVAL":
        raise HTTPException(
            status_code=400,
            detail=f"Batch is in status {batch.status!r} — must be AWAITING_APPROVAL to approve."
        )
    batch.status = "APPROVED"
    batch.approved_at = datetime.utcnow()
    await db.commit()
    await db.refresh(batch)
    return _serialize_batch(batch)


@router.post("/import/cancel/{batch_id}")
async def cancel_import(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_IMPORT_UPLOAD)),
):
    batch = await db.get(PHImportBatch, batch_id)
    if not batch:
        raise _not_found("Import batch", batch_id)
    if batch.status in ("COMPLETED", "IMPORTING"):
        raise HTTPException(
            status_code=400,
            detail=f"Batch is in status {batch.status!r} — cannot cancel an already completed or running import."
        )
    batch.status = "CANCELLED"
    await db.commit()
    await db.refresh(batch)
    return _serialize_batch(batch)


def _parse_excel_date(v: Any) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%b-%y", "%d-%b-%Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_excel_decimal(v: Any) -> Decimal:
    if v is None:
        return Decimal("0.00")
    s = str(v).strip().replace("$", "").replace(",", "")
    if not s or s.lower() == "none":
        return Decimal("0.00")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except Exception:
        return Decimal("0.00")


def _parse_excel_int(v: Any) -> int:
    if v is None:
        return 0
    s = str(v).strip().replace(",", "")
    try:
        return int(float(s))
    except Exception:
        return 0


@router.post("/import/analyze")
async def analyze_import(
    file: UploadFile = File(...),
    worksheet_name: str | None = Query(None),
    fiscal_year: str = Query("2026-2027"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission(Permissions.PH_IMPORT_UPLOAD)),
):
    """Step 2 & 3: Read uploaded workbook, validate rows, match employees & clients, and stage rows."""
    contents = await file.read()
    if not contents:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    try:
        wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {e}")

    # Determine which worksheet to use
    sheet_name = worksheet_name
    if not sheet_name or sheet_name not in wb.sheetnames:
        for candidate in ["Enhancement Programs", "Family Support", "Sheet1", "DATA (Dec-Mar)"]:
            if candidate in wb.sheetnames:
                sheet_name = candidate
                break
        if not sheet_name:
            sheet_name = wb.sheetnames[0]

    sheet = wb[sheet_name]

    # Detect record type
    sn_lower = sheet_name.lower()
    if "enhancement" in sn_lower:
        record_type = "ENHANCEMENT_PROGRAM"
        data_start_row = 8
    elif "family" in sn_lower:
        record_type = "FAMILY_SUPPORT"
        data_start_row = 8
    elif "other" in sn_lower or "eagle" in sn_lower or sn_lower == "sheet1":
        record_type = "OTHER_NATIONS"
        data_start_row = 4
    else:
        record_type = "ENHANCEMENT_PROGRAM"
        data_start_row = 8

    batch = PHImportBatch(
        uploaded_by_id=current_user.id,
        original_filename=file.filename or "upload.xlsx",
        worksheet_name=sheet_name,
        record_type=record_type,
        status="AWAITING_APPROVAL",
    )
    db.add(batch)
    await db.flush()

    total_rows = 0
    valid_rows = 0
    invalid_rows = 0
    employee_issues = set()
    client_issues = set()

    # Pre-load known employees, pending profiles, and client refs
    res_emps = await db.execute(select(Employee))
    known_employees = res_emps.scalars().all()
    emp_map = {f"{e.first_name} {e.last_name}".strip().lower(): e for e in known_employees}

    res_pending = await db.execute(select(PendingStaffProfile))
    known_pending = res_pending.scalars().all()
    pending_map = {p.display_name_from_excel.strip().lower(): p for p in known_pending}

    res_crefs = await db.execute(select(PHClientRef))
    known_crefs = res_crefs.scalars().all()
    cref_map = {cr.legacy_id.strip(): cr for cr in known_crefs}

    staged_rows = []

    for r in range(data_start_row, sheet.max_row + 1):
        if record_type == "ENHANCEMENT_PROGRAM":
            dt_raw = sheet.cell(r, 2).value
            emp_name = sheet.cell(r, 3).value
            pos_raw = sheet.cell(r, 4).value
            prog = sheet.cell(r, 5).value
            prog_act = sheet.cell(r, 6).value
            dt_prov = sheet.cell(r, 7).value
            city = sheet.cell(r, 8).value
            province = sheet.cell(r, 9).value
            partners = sheet.cell(r, 10).value
            closing_dt = sheet.cell(r, 11).value
            fam_in = sheet.cell(r, 12).value
            fam_out = sheet.cell(r, 13).value
            c_0_12 = sheet.cell(r, 16).value
            y_13_17 = sheet.cell(r, 17).value
            pm_18_26 = sheet.cell(r, 18).value
            p_parents = sheet.cell(r, 19).value
            p_members = sheet.cell(r, 20).value
            p_others = sheet.cell(r, 21).value
            acct = sheet.cell(r, 23).value
            sub = sheet.cell(r, 24).value
            exp_gc = sheet.cell(r, 25).value
            exp_po = sheet.cell(r, 26).value
            exp_eft = sheet.cell(r, 27).value
            exp_chq = sheet.cell(r, 28).value
            exp_cc = sheet.cell(r, 29).value
            exp_crt = sheet.cell(r, 30).value
            tot_cost_excel = sheet.cell(r, 31).value
            comments = sheet.cell(r, 32).value

            if not any([emp_name, prog, dt_raw, tot_cost_excel]):
                continue

            total_rows += 1
            errors = []
            rec_date = _parse_excel_date(dt_raw)
            if not prog or not str(prog).strip():
                errors.append("Missing required Program field.")

            e_gc = _parse_excel_decimal(exp_gc)
            e_po = _parse_excel_decimal(exp_po)
            e_eft = _parse_excel_decimal(exp_eft)
            e_chq = _parse_excel_decimal(exp_chq)
            e_cc = _parse_excel_decimal(exp_cc)
            e_crt = _parse_excel_decimal(exp_crt)
            calc_cost = e_gc + e_po + e_eft + e_chq + e_cc + e_crt

            if tot_cost_excel is not None:
                exp_stated = _parse_excel_decimal(tot_cost_excel)
                if abs(calc_cost - exp_stated) > Decimal("0.01") and exp_stated > Decimal("0.00") and calc_cost > Decimal("0.00"):
                    errors.append(f"Expense sum mismatch: calculated ${calc_cost} vs stated ${exp_stated}")

            emp_name_str = str(emp_name).strip() if emp_name else None
            resolved_emp_id = None
            resolved_pending_id = None
            emp_match_status = "NO_MATCH"

            if emp_name_str:
                emp_key = emp_name_str.lower()
                if emp_key in emp_map:
                    resolved_emp_id = emp_map[emp_key].id
                    emp_match_status = "MATCHED"
                elif emp_key in pending_map:
                    resolved_pending_id = pending_map[emp_key].id
                    emp_match_status = "PENDING_PROFILE"
                else:
                    parts = emp_name_str.split(" ", 1)
                    fn = parts[0]
                    ln = parts[1] if len(parts) > 1 else ""
                    new_pending = PendingStaffProfile(
                        first_name=fn,
                        last_name=ln,
                        display_name_from_excel=emp_name_str,
                        position=str(pos_raw).strip() if pos_raw else "Prevention Worker",
                        department="Prevention & Healing",
                        onboarding_status="PENDING",
                        source_batch_id=batch.id,
                    )
                    db.add(new_pending)
                    await db.flush()
                    pending_map[emp_key] = new_pending
                    resolved_pending_id = new_pending.id
                    emp_match_status = "PENDING_PROFILE_CREATED"
                    employee_issues.add(emp_name_str)

            row_dict = {
                "record_date": rec_date.isoformat() if rec_date else None,
                "employee_name_raw": emp_name_str,
                "position_raw": str(pos_raw).strip() if pos_raw else None,
                "program": str(prog).strip() if prog else "General Prevention",
                "program_activity": str(prog_act).strip() if prog_act else None,
                "date_provided": _parse_excel_date(dt_prov).isoformat() if _parse_excel_date(dt_prov) else None,
                "closing_date": _parse_excel_date(closing_dt).isoformat() if _parse_excel_date(closing_dt) else None,
                "city": str(city).strip() if city else None,
                "province": str(province).strip() if province else None,
                "partners": str(partners).strip() if partners else None,
                "families_in_province": _parse_excel_int(fam_in),
                "families_out_of_province": _parse_excel_int(fam_out),
                "participants_children_0_12": _parse_excel_int(c_0_12),
                "participants_youth_13_17": _parse_excel_int(y_13_17),
                "participants_post_majority_18_26": _parse_excel_int(pm_18_26),
                "participants_parents_caregivers": _parse_excel_int(p_parents),
                "participants_family_members": _parse_excel_int(p_members),
                "participants_others": _parse_excel_int(p_others),
                "account_code": str(acct).strip() if acct else None,
                "subcode": str(sub).strip() if sub else None,
                "expense_gift_card": float(e_gc),
                "expense_purchase_order": float(e_po),
                "expense_eft": float(e_eft),
                "expense_cheque": float(e_chq),
                "expense_credit_card": float(e_cc),
                "expense_courtesy_card": float(e_crt),
                "total_cost": float(calc_cost if calc_cost > 0 else _parse_excel_decimal(tot_cost_excel)),
                "comments": str(comments).strip() if comments else None,
                "fiscal_year": fiscal_year,
            }

            status_val = "VALID" if not errors else "INVALID"
            if status_val == "VALID":
                valid_rows += 1
            else:
                invalid_rows += 1

            im_row = PHImportRow(
                batch_id=batch.id,
                source_row_number=r,
                status=status_val,
                validation_errors="; ".join(errors) if errors else None,
                employee_name_raw=emp_name_str,
                employee_match_status=emp_match_status,
                resolved_employee_id=resolved_emp_id,
                resolved_pending_staff_id=resolved_pending_id,
                row_data=row_dict,
            )
            db.add(im_row)
            staged_rows.append(im_row)

        elif record_type == "FAMILY_SUPPORT":
            dt_raw = sheet.cell(r, 2).value
            emp_name = sheet.cell(r, 3).value
            pos_raw = sheet.cell(r, 4).value
            service_area = sheet.cell(r, 5).value
            prog = sheet.cell(r, 6).value
            st_type = sheet.cell(r, 7).value
            description = sheet.cell(r, 8).value
            partners = sheet.cell(r, 9).value
            dt_prov = sheet.cell(r, 10).value
            rm_id = sheet.cell(r, 11).value
            city = sheet.cell(r, 12).value
            province = sheet.cell(r, 13).value
            c_0_12 = sheet.cell(r, 14).value
            y_13_17 = sheet.cell(r, 15).value
            pm_18_26 = sheet.cell(r, 16).value
            a_18_plus = sheet.cell(r, 17).value
            acct = sheet.cell(r, 18).value
            sub = sheet.cell(r, 19).value
            proj = sheet.cell(r, 20).value
            exp_gc = sheet.cell(r, 21).value
            exp_po = sheet.cell(r, 22).value
            exp_eft = sheet.cell(r, 23).value
            exp_chq = sheet.cell(r, 24).value
            exp_cc = sheet.cell(r, 25).value
            exp_petty = sheet.cell(r, 26).value
            exp_crt = sheet.cell(r, 27).value
            tot_cost_excel = sheet.cell(r, 28).value
            comments = sheet.cell(r, 29).value

            if not any([emp_name, prog, dt_raw, rm_id, tot_cost_excel]):
                continue

            total_rows += 1
            errors = []
            rec_date = _parse_excel_date(dt_raw)
            if not prog or not str(prog).strip():
                errors.append("Missing required Program field.")

            e_gc = _parse_excel_decimal(exp_gc)
            e_po = _parse_excel_decimal(exp_po)
            e_eft = _parse_excel_decimal(exp_eft)
            e_chq = _parse_excel_decimal(exp_chq)
            e_cc = _parse_excel_decimal(exp_cc)
            e_petty = _parse_excel_decimal(exp_petty)
            e_crt = _parse_excel_decimal(exp_crt)
            calc_cost = e_gc + e_po + e_eft + e_chq + e_cc + e_petty + e_crt

            if tot_cost_excel is not None:
                exp_stated = _parse_excel_decimal(tot_cost_excel)
                if abs(calc_cost - exp_stated) > Decimal("0.01") and exp_stated > Decimal("0.00") and calc_cost > Decimal("0.00"):
                    errors.append(f"Expense sum mismatch: calculated ${calc_cost} vs stated ${exp_stated}")

            emp_name_str = str(emp_name).strip() if emp_name else None
            resolved_emp_id = None
            resolved_pending_id = None
            emp_match_status = "NO_MATCH"

            if emp_name_str:
                emp_key = emp_name_str.lower()
                if emp_key in emp_map:
                    resolved_emp_id = emp_map[emp_key].id
                    emp_match_status = "MATCHED"
                elif emp_key in pending_map:
                    resolved_pending_id = pending_map[emp_key].id
                    emp_match_status = "PENDING_PROFILE"
                else:
                    parts = emp_name_str.split(" ", 1)
                    fn = parts[0]
                    ln = parts[1] if len(parts) > 1 else ""
                    new_pending = PendingStaffProfile(
                        first_name=fn,
                        last_name=ln,
                        display_name_from_excel=emp_name_str,
                        position=str(pos_raw).strip() if pos_raw else "Prevention Worker",
                        department="Prevention & Healing",
                        onboarding_status="PENDING",
                        source_batch_id=batch.id,
                    )
                    db.add(new_pending)
                    await db.flush()
                    pending_map[emp_key] = new_pending
                    resolved_pending_id = new_pending.id
                    emp_match_status = "PENDING_PROFILE_CREATED"
                    employee_issues.add(emp_name_str)

            # Client matching (RedMane ID)
            rm_id_str = str(rm_id).strip() if rm_id is not None and str(rm_id).strip() != "" else None
            resolved_pid = None
            client_match_status = "NO_MATCH"

            if rm_id_str:
                if rm_id_str in cref_map:
                    cref = cref_map[rm_id_str]
                    if cref.status == "CONFIRMED" and cref.person_id:
                        resolved_pid = cref.person_id
                        client_match_status = "CONFIRMED"
                    else:
                        client_match_status = "UNRESOLVED"
                        client_issues.add(rm_id_str)
                else:
                    new_cref = PHClientRef(
                        legacy_id=rm_id_str,
                        legacy_id_type="REDMANE",
                        status="UNRESOLVED",
                    )
                    db.add(new_cref)
                    await db.flush()
                    cref_map[rm_id_str] = new_cref
                    client_match_status = "UNRESOLVED"
                    client_issues.add(rm_id_str)

            row_dict = {
                "record_date": rec_date.isoformat() if rec_date else None,
                "employee_name_raw": emp_name_str,
                "position_raw": str(pos_raw).strip() if pos_raw else None,
                "service_area": str(service_area).strip() if service_area else None,
                "program": str(prog).strip() if prog else "Family Support",
                "support_type": str(st_type).strip() if st_type else None,
                "description": str(description).strip() if description else None,
                "partners": str(partners).strip() if partners else None,
                "date_provided": _parse_excel_date(dt_prov).isoformat() if _parse_excel_date(dt_prov) else None,
                "client_ref_legacy_id": rm_id_str,
                "city": str(city).strip() if city else None,
                "province": str(province).strip() if province else None,
                "participants_children_0_12": _parse_excel_int(c_0_12),
                "participants_youth_13_17": _parse_excel_int(y_13_17),
                "participants_post_majority_18_26": _parse_excel_int(pm_18_26),
                "participants_adults_18_plus": _parse_excel_int(a_18_plus),
                "account_code": str(acct).strip() if acct else None,
                "subcode": str(sub).strip() if sub else None,
                "project_code": str(proj).strip() if proj else None,
                "expense_gift_card": float(e_gc),
                "expense_purchase_order": float(e_po),
                "expense_eft": float(e_eft),
                "expense_cheque": float(e_chq),
                "expense_credit_card": float(e_cc),
                "expense_petty_cash": float(e_petty),
                "expense_courtesy_card": float(e_crt),
                "total_cost": float(calc_cost if calc_cost > 0 else _parse_excel_decimal(tot_cost_excel)),
                "comments": str(comments).strip() if comments else None,
                "fiscal_year": fiscal_year,
                "client_match_status": client_match_status,
            }

            status_val = "VALID" if not errors else "INVALID"
            if status_val == "VALID":
                valid_rows += 1
            else:
                invalid_rows += 1

            im_row = PHImportRow(
                batch_id=batch.id,
                source_row_number=r,
                status=status_val,
                validation_errors="; ".join(errors) if errors else None,
                employee_name_raw=emp_name_str,
                employee_match_status=emp_match_status,
                resolved_employee_id=resolved_emp_id,
                resolved_pending_staff_id=resolved_pending_id,
                client_ref_raw=rm_id_str,
                client_match_status=client_match_status,
                resolved_person_id=resolved_pid,
                row_data=row_dict,
            )
            db.add(im_row)
            staged_rows.append(im_row)

        else:  # OTHER_NATIONS
            dt_raw = sheet.cell(r, 1).value
            worker_name = sheet.cell(r, 2).value
            fam_ref = sheet.cell(r, 3).value
            serv_type = sheet.cell(r, 4).value
            cost_raw = sheet.cell(r, 5).value
            status_num = sheet.cell(r, 6).value
            cowessess_conn = sheet.cell(r, 7).value
            non_cowessess = sheet.cell(r, 8).value
            on_res_raw = sheet.cell(r, 9).value
            off_res_raw = sheet.cell(r, 10).value
            comments = sheet.cell(r, 11).value

            if not any([worker_name, fam_ref, serv_type, cost_raw, dt_raw]):
                continue

            total_rows += 1
            errors = []
            rec_date = _parse_excel_date(dt_raw)
            cost_dec = _parse_excel_decimal(cost_raw)

            on_res = None
            if on_res_raw is not None and str(on_res_raw).strip():
                on_res = "yes" in str(on_res_raw).lower()
            elif off_res_raw is not None and str(off_res_raw).strip():
                on_res = False

            worker_str = str(worker_name).strip() if worker_name else None
            resolved_emp_id = None
            resolved_pending_id = None
            emp_match_status = "NO_MATCH"

            if worker_str:
                emp_key = worker_str.lower()
                if emp_key in emp_map:
                    resolved_emp_id = emp_map[emp_key].id
                    emp_match_status = "MATCHED"
                elif emp_key in pending_map:
                    resolved_pending_id = pending_map[emp_key].id
                    emp_match_status = "PENDING_PROFILE"
                else:
                    parts = worker_str.split(" ", 1)
                    fn = parts[0]
                    ln = parts[1] if len(parts) > 1 else ""
                    new_pending = PendingStaffProfile(
                        first_name=fn,
                        last_name=ln,
                        display_name_from_excel=worker_str,
                        position="Prevention Worker",
                        department="Prevention & Healing",
                        onboarding_status="PENDING",
                        source_batch_id=batch.id,
                    )
                    db.add(new_pending)
                    await db.flush()
                    pending_map[emp_key] = new_pending
                    resolved_pending_id = new_pending.id
                    emp_match_status = "PENDING_PROFILE_CREATED"
                    employee_issues.add(worker_str)

            row_dict = {
                "record_date": rec_date.isoformat() if rec_date else None,
                "worker_name_raw": worker_str,
                "family_reference_raw": str(fam_ref).strip() if fam_ref else None,
                "service_type": str(serv_type).strip() if serv_type else "Support",
                "cost": float(cost_dec),
                "on_reserve": on_res,
                "cowessess_connection": str(cowessess_conn).strip() if cowessess_conn else None,
                "other_community": str(non_cowessess).strip() if non_cowessess else (str(off_res_raw).strip() if off_res_raw else None),
                "historical_reference_number": str(status_num).strip() if status_num else None,
                "comments": str(comments).strip() if comments else None,
                "fiscal_year": fiscal_year,
            }

            status_val = "VALID" if not errors else "INVALID"
            if status_val == "VALID":
                valid_rows += 1
            else:
                invalid_rows += 1

            im_row = PHImportRow(
                batch_id=batch.id,
                source_row_number=r,
                status=status_val,
                validation_errors="; ".join(errors) if errors else None,
                employee_name_raw=worker_str,
                employee_match_status=emp_match_status,
                resolved_employee_id=resolved_emp_id,
                resolved_pending_staff_id=resolved_pending_id,
                client_ref_raw=str(fam_ref).strip() if fam_ref else None,
                row_data=row_dict,
            )
            db.add(im_row)
            staged_rows.append(im_row)

    batch.total_rows_analyzed = total_rows
    batch.valid_rows = valid_rows
    batch.invalid_rows = invalid_rows
    batch.status = "AWAITING_APPROVAL"

    await db.commit()
    await db.refresh(batch)

    return {
        "batch_id": str(batch.id),
        "filename": batch.original_filename,
        "worksheet_name": batch.worksheet_name,
        "record_type": batch.record_type,
        "total_rows_analyzed": total_rows,
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "employee_matching_issues": sorted(list(employee_issues)),
        "client_matching_issues": sorted(list(client_issues)),
        "preview_rows": [_serialize_row(r) for r in staged_rows[:25]],
        "ready_for_approval": valid_rows > 0,
    }


@router.post("/import/execute/{batch_id}")
async def execute_import(
    batch_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_permission(Permissions.PH_IMPORT_EXECUTE)),
):
    """Step 8: Execute approved import in a single atomic transaction."""
    batch = await db.get(PHImportBatch, batch_id)
    if not batch:
        raise _not_found("Import batch", batch_id)
    if batch.status != "APPROVED":
        raise HTTPException(
            status_code=400,
            detail=f"Batch is in status {batch.status!r} — must be APPROVED before execution."
        )

    q = select(PHImportRow).where(PHImportRow.batch_id == batch_id, PHImportRow.status == "VALID")
    res = await db.execute(q)
    rows = res.scalars().all()
    if not rows:
        raise HTTPException(status_code=400, detail="No VALID rows found in this batch to import.")

    imported_count = 0
    for r in rows:
        data = r.row_data or {}
        if batch.record_type == "ENHANCEMENT_PROGRAM":
            rec = EnhancementProgram(
                employee_id=r.resolved_employee_id,
                pending_staff_id=r.resolved_pending_staff_id,
                employee_name_raw=r.employee_name_raw,
                position_raw=data.get("position_raw"),
                record_date=_parse_excel_date(data.get("record_date")),
                program=data.get("program") or "General Program",
                program_activity=data.get("program_activity"),
                date_provided=_parse_excel_date(data.get("date_provided")),
                closing_date=_parse_excel_date(data.get("closing_date")),
                city=data.get("city"),
                province=data.get("province"),
                partners=data.get("partners"),
                families_in_province=int(data.get("families_in_province") or 0),
                families_out_of_province=int(data.get("families_out_of_province") or 0),
                participants_children_0_12=int(data.get("participants_children_0_12") or 0),
                participants_youth_13_17=int(data.get("participants_youth_13_17") or 0),
                participants_post_majority_18_26=int(data.get("participants_post_majority_18_26") or 0),
                participants_parents_caregivers=int(data.get("participants_parents_caregivers") or 0),
                participants_family_members=int(data.get("participants_family_members") or 0),
                participants_others=int(data.get("participants_others") or 0),
                account_code=data.get("account_code"),
                subcode=data.get("subcode"),
                expense_gift_card=_parse_excel_decimal(data.get("expense_gift_card")),
                expense_purchase_order=_parse_excel_decimal(data.get("expense_purchase_order")),
                expense_eft=_parse_excel_decimal(data.get("expense_eft")),
                expense_cheque=_parse_excel_decimal(data.get("expense_cheque")),
                expense_credit_card=_parse_excel_decimal(data.get("expense_credit_card")),
                expense_courtesy_card=_parse_excel_decimal(data.get("expense_courtesy_card")),
                comments=data.get("comments"),
                fiscal_year=data.get("fiscal_year") or "2026-2027",
                import_batch_id=batch.id,
                source_row_number=r.source_row_number,
            )
            db.add(rec)
            await db.flush()
            r.status = "IMPORTED"
            r.imported_record_id = rec.id
            r.imported_record_type = "ENHANCEMENT_PROGRAM"
            imported_count += 1

        elif batch.record_type == "FAMILY_SUPPORT":
            rec = FamilySupportRecord(
                employee_id=r.resolved_employee_id,
                pending_staff_id=r.resolved_pending_staff_id,
                employee_name_raw=r.employee_name_raw,
                position_raw=data.get("position_raw"),
                service_area=data.get("service_area"),
                record_date=_parse_excel_date(data.get("record_date")),
                program=data.get("program") or "Family Support",
                support_type=data.get("support_type"),
                description=data.get("description"),
                partners=data.get("partners"),
                date_provided=_parse_excel_date(data.get("date_provided")),
                person_id=r.resolved_person_id,
                client_ref_legacy_id=r.client_ref_raw,
                city=data.get("city"),
                province=data.get("province"),
                participants_children_0_12=int(data.get("participants_children_0_12") or 0),
                participants_youth_13_17=int(data.get("participants_youth_13_17") or 0),
                participants_post_majority_18_26=int(data.get("participants_post_majority_18_26") or 0),
                participants_adults_18_plus=int(data.get("participants_adults_18_plus") or 0),
                account_code=data.get("account_code"),
                subcode=data.get("subcode"),
                project_code=data.get("project_code"),
                expense_gift_card=_parse_excel_decimal(data.get("expense_gift_card")),
                expense_purchase_order=_parse_excel_decimal(data.get("expense_purchase_order")),
                expense_eft=_parse_excel_decimal(data.get("expense_eft")),
                expense_cheque=_parse_excel_decimal(data.get("expense_cheque")),
                expense_credit_card=_parse_excel_decimal(data.get("expense_credit_card")),
                expense_petty_cash=_parse_excel_decimal(data.get("expense_petty_cash")),
                expense_courtesy_card=_parse_excel_decimal(data.get("expense_courtesy_card")),
                comments=data.get("comments"),
                fiscal_year=data.get("fiscal_year") or "2026-2027",
                client_match_status=r.client_match_status,
                import_batch_id=batch.id,
                source_row_number=r.source_row_number,
            )
            db.add(rec)
            await db.flush()
            r.status = "IMPORTED"
            r.imported_record_id = rec.id
            r.imported_record_type = "FAMILY_SUPPORT"
            imported_count += 1

        else:  # OTHER_NATIONS
            rec = OtherNationsSupport(
                employee_id=r.resolved_employee_id,
                pending_staff_id=r.resolved_pending_staff_id,
                worker_name_raw=r.employee_name_raw,
                record_date=_parse_excel_date(data.get("record_date")),
                family_reference_raw=data.get("family_reference_raw"),
                service_type=data.get("service_type") or "Support",
                cost=_parse_excel_decimal(data.get("cost")),
                on_reserve=data.get("on_reserve"),
                cowessess_connection=data.get("cowessess_connection"),
                other_community=data.get("other_community"),
                historical_reference_number=data.get("historical_reference_number"),
                comments=data.get("comments"),
                person_id=r.resolved_person_id,
                client_match_status=r.client_match_status,
                fiscal_year=data.get("fiscal_year") or "2026-2027",
                import_batch_id=batch.id,
                source_row_number=r.source_row_number,
            )
            db.add(rec)
            await db.flush()
            r.status = "IMPORTED"
            r.imported_record_id = rec.id
            r.imported_record_type = "OTHER_NATIONS"
            imported_count += 1

    batch.status = "COMPLETED"
    batch.completed_at = datetime.utcnow()
    batch.imported_rows = imported_count

    await db.commit()
    await db.refresh(batch)
    return _serialize_batch(batch)


def _serialize_batch(r: PHImportBatch) -> dict:
    return {
        "id": str(r.id),
        "uploaded_by_id": str(r.uploaded_by_id),
        "original_filename": r.original_filename,
        "worksheet_name": r.worksheet_name,
        "record_type": r.record_type,
        "total_rows_analyzed": r.total_rows_analyzed,
        "valid_rows": r.valid_rows,
        "invalid_rows": r.invalid_rows,
        "imported_rows": r.imported_rows,
        "status": r.status,
        "approved_by_id": str(r.approved_by_id) if r.approved_by_id else None,
        "approved_at": r.approved_at.isoformat() if r.approved_at else None,
        "completed_at": r.completed_at.isoformat() if r.completed_at else None,
        "error_summary": r.error_summary,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


def _serialize_row(r: PHImportRow) -> dict:
    return {
        "id": str(r.id),
        "batch_id": str(r.batch_id),
        "source_row_number": r.source_row_number,
        "status": r.status,
        "validation_errors": r.validation_errors,
        "imported_record_id": str(r.imported_record_id) if r.imported_record_id else None,
        "imported_record_type": r.imported_record_type,
        "employee_name_raw": r.employee_name_raw,
        "employee_match_status": r.employee_match_status,
        "resolved_employee_id": str(r.resolved_employee_id) if r.resolved_employee_id else None,
        "client_ref_raw": r.client_ref_raw,
        "client_match_status": r.client_match_status,
        "resolved_person_id": str(r.resolved_person_id) if r.resolved_person_id else None,
        "row_data": r.row_data,
    }


# ──────────────────────────────────────────────────────────────
# Annual Summary / Reports
# ──────────────────────────────────────────────────────────────

@router.get("/summary")
async def ph_summary(
    fiscal_year: str = Query("2026-2027"),
    db: AsyncSession = Depends(get_db),
    _: Any = Depends(require_permission(Permissions.PH_REPORT_READ)),
):
    """Return aggregated annual summary statistics for the Prevention & Healing workspace."""

    # Enhancement Programs totals
    ep_q = await db.execute(
        select(
            func.count(EnhancementProgram.id).label("ep_count"),
            func.coalesce(
                func.sum(
                    EnhancementProgram.expense_gift_card +
                    EnhancementProgram.expense_purchase_order +
                    EnhancementProgram.expense_eft +
                    EnhancementProgram.expense_cheque +
                    EnhancementProgram.expense_credit_card +
                    EnhancementProgram.expense_courtesy_card
                ), 0
            ).label("ep_total_cost"),
            func.coalesce(func.sum(EnhancementProgram.families_in_province + EnhancementProgram.families_out_of_province), 0).label("ep_total_families"),
            func.coalesce(
                func.sum(
                    EnhancementProgram.participants_children_0_12 +
                    EnhancementProgram.participants_youth_13_17 +
                    EnhancementProgram.participants_post_majority_18_26 +
                    EnhancementProgram.participants_parents_caregivers +
                    EnhancementProgram.participants_family_members +
                    EnhancementProgram.participants_others
                ), 0
            ).label("ep_total_participants"),
        ).where(EnhancementProgram.fiscal_year == fiscal_year)
    )
    ep = ep_q.one()

    # Family Support totals
    fsr_q = await db.execute(
        select(
            func.count(FamilySupportRecord.id).label("fsr_count"),
            func.coalesce(
                func.sum(
                    FamilySupportRecord.expense_gift_card +
                    FamilySupportRecord.expense_purchase_order +
                    FamilySupportRecord.expense_eft +
                    FamilySupportRecord.expense_cheque +
                    FamilySupportRecord.expense_credit_card +
                    FamilySupportRecord.expense_petty_cash +
                    FamilySupportRecord.expense_courtesy_card
                ), 0
            ).label("fsr_total_cost"),
        ).where(FamilySupportRecord.fiscal_year == fiscal_year)
    )
    fsr = fsr_q.one()

    # Other Nations totals
    ons_q = await db.execute(
        select(
            func.count(OtherNationsSupport.id).label("ons_count"),
            func.coalesce(func.sum(OtherNationsSupport.cost), 0).label("ons_total_cost"),
        ).where(OtherNationsSupport.fiscal_year == fiscal_year)
    )
    ons = ons_q.one()

    # Pending staff
    ps_q = await db.execute(
        select(func.count(PendingStaffProfile.id)).where(
            PendingStaffProfile.onboarding_status.in_(["PENDING", "EMAIL_NEEDED"])
        )
    )
    pending_staff_count = ps_q.scalar() or 0

    # Unresolved client refs
    cr_q = await db.execute(
        select(func.count(PHClientRef.id)).where(PHClientRef.status == "UNRESOLVED")
    )
    unresolved_client_refs = cr_q.scalar() or 0

    return {
        "fiscal_year": fiscal_year,
        "enhancement_programs": {
            "record_count": ep.ep_count,
            "total_cost": float(ep.ep_total_cost),
            "total_families": int(ep.ep_total_families),
            "total_participants": int(ep.ep_total_participants),
        },
        "family_support": {
            "record_count": fsr.fsr_count,
            "total_cost": float(fsr.fsr_total_cost),
        },
        "other_nations_support": {
            "record_count": ons.ons_count,
            "total_cost": float(ons.ons_total_cost),
        },
        "combined_total_cost": float(
            (ep.ep_total_cost or 0) + (fsr.fsr_total_cost or 0) + (ons.ons_total_cost or 0)
        ),
        "pending_staff_needing_action": pending_staff_count,
        "unresolved_client_refs": unresolved_client_refs,
    }
