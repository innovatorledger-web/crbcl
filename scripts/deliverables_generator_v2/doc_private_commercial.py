"""Generate Private Developer Commercial Correspondence (DOCX & PDF).

Dedicated Folder:
deliverables/CRBCL_Private_Developer_Commercial_Correspondence_October_2026/

Files:
CRBCL_Private_Original_Agreement_Additional_Development_and_Compensation_Discussion_October_2026.docx
CRBCL_Private_Original_Agreement_Additional_Development_and_Compensation_Discussion_October_2026.pdf

Audience: Strictly the individual who originally hired the developer (NOT the CRBCL Board).
Tone: Professional, respectful, factual, first-person perspective.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib import colors

from scripts.deliverables_generator_v2.common_v2 import (
    PRIVATE_COMMERCIAL_PACKAGE_DIR, NumberedCanvasPrivate, get_pdf_styles_v2,
    PRIMARY_HEX, SECONDARY_HEX, DARK_HEX, LIGHT_BG_HEX, BORDER_HEX,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_DARK, COLOR_LIGHT_BG, COLOR_BORDER,
    set_cell_background, set_cell_margins
)

DOC_NAME_BASE = "CRBCL_Private_Original_Agreement_Additional_Development_and_Compensation_Discussion_October_2026"

COMMERCIAL_FACTS_TABLE = [
    ("Original Project Quotation", "CAD $2,500.00", "Initial agreement based on the preliminary, more limited scope discussed at project outset (basic case notes, simple client index)."),
    ("Payments Received to Date", "CAD $2,200.00", "Disbursed in progress payments during initial platform setup."),
    ("Nominal Difference on Original Quote", "CAD $300.00", "Unpaid nominal difference between original quotation and payments received to date; noted factually for reconciliation."),
    ("Formally Agreed Additional Compensation", "CAD $0.00 (Pending Discussion)", "Verbal indication provided that extra compensation would be arranged; no formal dollar figure or change order signed to date."),
    ("Scope Expansion Ratio", "Approximately 12x to 15x", "Evolution from a basic ~15-table prototype into an enterprise platform with 205 mapped models, 521 endpoints, 35 admin workspaces, 24 team dashboards, and a mobile application."),
]

EXPANDED_WORK_ITEMS = [
    (
        "1. Departmental Dashboards (24 Dedicated Teams)",
        "Original discussion assumed a single shared case view.",
        "Architected and implemented 24 dedicated Team Dashboards (Executive, CFS, Prevention & Healing, Sacred Wolf Lodge, Home Fire, Family Support, Clinical, Intake, Legal, HR, Finance, IT, etc.) with customized work queues, KPI aggregations, and role-based routing.",
        "Implemented in source (src/pages/Teams.jsx, TeamDashboard.jsx, IDs 1-24); authenticated testing conducted on CEO and Board dashboards."
    ),
    (
        "2. Role-Based Access Control & Permissions Engine",
        "Initial concept assumed basic binary user permissions ('admin' vs 'caseworker').",
        "Engineered an enterprise RBAC engine with 327 discrete permissions across 19 canonical system roles, seeded in database migration 030, with granular endpoint authorization decorators.",
        "Fully implemented in backend/app/permissions/constants.py and verified via comprehensive pytest suites."
    ),
    (
        "3. Administrative Control Centre (35 Workspaces)",
        "No dedicated administrative hub planned beyond standard user management.",
        "Created 35 canonical administrative workspaces covering fleet management, facility reservations, lookups, security, and immutable audit logs.",
        "Implemented in source (backend/app/services/dashboard_registry_service.py CANONICAL_WORKSPACES)."
    ),
    (
        "4. Executive & Board Governance Portals",
        "Not included in original specification.",
        "Engineered dedicated CEO Executive Portal (/ceo) and Board of Governors Dashboard (/board) with high-level KPI visual summaries and formal report publication workflows.",
        "Implemented in source and live verified during September 2026 authenticated sessions (commits 0319f07 and d44e3c9)."
    ),
    (
        "5. Case Conflict-of-Interest Blinding Engine",
        "Not identified in original quotation requirements.",
        "Built dynamic Case Restriction engine that actively filters and blinds conflicted caseworkers on all case and note endpoints to uphold community privacy and customary care ethics.",
        "Implemented in backend/app/services/case_restriction_service.py and database migration 006."
    ),
    (
        "6. Canonical Master Person Index & Client Records",
        "Basic flat client contact table.",
        "Architected decoupled Master Person Index (MPI) with deduplication, alias tracking, family relationship graphs, and statutory child protection lifecycles.",
        "Implemented across backend/app/models/ (clients, persons, family) and migration series."
    ),
    (
        "7. Employee HR Management & Login Quarantine",
        "Direct user account creation without separation.",
        "Decoupled employee HR records from platform user accounts and implemented login quarantine security preventing unapproved staff access until HR verification.",
        "Implemented in backend/app/models/employee.py, migration 017, and test_employee_quarantine.py."
    ),
    (
        "8. Prevention & Healing Module (5 Tabs, 36 Endpoints)",
        "Not in original specification; added as a major initiative during platform evolution.",
        "Developed complete 5-tab workspace, 3-tier duplicate protection engine, safe Excel import wizard, financial coding, and 36 dedicated API endpoints.",
        "Implemented in source (backend/app/api/v1/ph/, PreventionHealing.jsx, Migration 031); verified in local test suite."
    ),
    (
        "9. Historical Spreadsheet Import Preparation",
        "Manual data re-entry assumed at outset.",
        "Engineered synthetic data test harness, legacy schema mapping, duplicate reconciliation workflows, and dry-run import pipelines for legacy Cowessess spreadsheets.",
        "Implemented and verified in synthetic testing; live community data import strictly pending CRBCL operational sign-off."
    ),
    (
        "10. Financial Reporting & Billing Workflows",
        "Simple expense tracking or manual bookkeeping.",
        "Engineered Jordan's Principle federal billing EDI claim tracking engine, service invoice logs, per diem billing calculators, and multi-fund ledger accounting models.",
        "Implemented in source across finance models, routers, and database migrations."
    ),
    (
        "11. Cross-Platform Flutter Mobile Application",
        "Web browser only on mobile devices.",
        "Developed native cross-platform Flutter application targeting Android, iOS, and Web, with dedicated screens for client lookup, intake drafting, visit notes, and dashboard navigation.",
        "Implemented in mobile/ codebase; updated Android APK compiled."
    ),
    (
        "12. Mobile Offline Functionality & SQLite Storage",
        "No offline mobile storage in original scope.",
        "Implemented encrypted local SQLite database (SQLCipher AES-256) with dynamic key storage in FlutterSecureStorage and transactional outbox queue for remote on-reserve fieldwork. Personally tested on physical Android device with confirmed offline operation.",
        "User tested and confirmed working by developer on physical Android device; formal organizational UAT pending."
    ),
    (
        "13. Isolated Website Offline Workspace",
        "Standard online-only web application.",
        "Implemented isolated browser offline workspace (/offline-workspace) using HTML5 Service Worker and IndexedDB, gated by environment variable CRBCL_OFFLINE_WEB_ENABLED.",
        "Implemented in source (commit e33c87a); production feature gate set to disabled pending multi-tab testing."
    ),
    (
        "14. Speech-to-Text Voice Dictation Pipeline",
        "Manual keyboard text entry only.",
        "Engineered ephemeral in-memory Whisper voice transcription gateway enabling frontline workers to dictate case notes and voice queries with zero audio disk retention.",
        "Implemented and live verified on website as an ephemeral input aid (POST /api/v1/ask-red-bear/transcribe)."
    ),
    (
        "15. Ask Red Bear Assistant Subsystem",
        "Not envisioned in original quotation.",
        "Built responsive frontend chat modal, streaming transcription gateway, and modular AiGateway abstraction layer with simulation stub (FakeAiProvider) ready for cloud LLM injection.",
        "Frontend and transcription live verified; AI answer generation stubbed pending Board selection of cloud provider."
    ),
    (
        "16. Enterprise Security, Privacy & Cryptography",
        "Basic password authentication.",
        "Implemented Argon2id password hashing, frontline 6-digit PIN login with brute-force lockout, SHA-256 ingestion provenance hashing, CSRF token defense, and CORS origin enforcement.",
        "Implemented in auth routers, security middleware, and verified in automated security test suites."
    ),
    (
        "17. Automated Testing, Debugging & Quality Assurance",
        "Informal testing of basic prototype.",
        "Constructed over 30 pytest automated test suites covering authentication, RBAC, clinical cases, sync logic, and duplicate prevention, using synthetic datasets.",
        "Active test suites passing in repository test runs."
    ),
    (
        "18. Comprehensive Technical Architecture Documentation",
        "Standard brief README file.",
        "Authored exhaustive architectural guides, API catalogs, schema migration histories, and governance audit reports to ensure complete technical transparency for leadership.",
        "Delivered in Revision 1 and Revision 2 documentation packages."
    ),
]


def build_private_commercial_doc():
    docx_path = os.path.join(PRIVATE_COMMERCIAL_PACKAGE_DIR, f"{DOC_NAME_BASE}.docx")
    pdf_path = os.path.join(PRIVATE_COMMERCIAL_PACKAGE_DIR, f"{DOC_NAME_BASE}.pdf")

    print(f"Building Private Commercial Document...")
    print(f"Target DOCX: {docx_path}")
    print(f"Target PDF:  {pdf_path}")

    # =========================================================================
    # BUILD DOCX
    # =========================================================================
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header and Footer
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = "PRIVATE & CONFIDENTIAL — FOR THE ORIGINAL HIRING CONTACT ONLY"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.style.font.name = "Arial"
    hp.style.font.size = Pt(8)
    hp.style.font.color.rgb = RGBColor(180, 83, 9)

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = "CONFIDENTIAL COMMERCIAL CORRESPONDENCE — NOT FOR BOARD DISTRIBUTION"
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.style.font.name = "Arial"
    fp.style.font.size = Pt(8)
    fp.style.font.color.rgb = RGBColor(100, 116, 139)

    # Notice Banner
    notice_tbl = doc.add_table(rows=1, cols=1)
    notice_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = notice_tbl.cell(0, 0)
    set_cell_background(c, "FEF3C7")
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    np = c.paragraphs[0]
    nh = np.add_run("PRIVATE AND CONFIDENTIAL — FOR THE ORIGINAL HIRING CONTACT ONLY\n\n")
    nh.bold = True
    nh.font.name = "Arial"
    nh.font.size = Pt(10)
    nh.font.color.rgb = RGBColor(146, 64, 14)
    nb = np.add_run(
        "This correspondence is strictly private and commercial in nature. It is addressed exclusively to the "
        "individual who engaged the developer's services for the Chief Red Bear Children's Lodge platform. "
        "It must not be circulated to the CRBCL Board of Directors, executive committees, or external evaluators. "
        "Technical and progress documentation for the Board has been provided separately in dedicated Board reports."
    )
    nb.font.name = "Arial"
    nb.font.size = Pt(9)
    nb.font.color.rgb = RGBColor(120, 53, 15)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("CHIEF RED BEAR CHILDREN'S LODGE (CRBCL) PLATFORM\nCOMMERCIAL REVIEW & COMPENSATION DISCUSSION")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(107, 29, 47)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Original Agreement, Scope Expansion, Completed Engineering, and Outstanding Compensation Discussion\nPrivate Correspondence — October 2026")
    r_sub.bold = True
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(201, 138, 44)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    def add_h1(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(107, 29, 47)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(9.5)
            r_pre.font.color.rgb = RGBColor(30, 41, 59)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_after = Pt(4.5)
        return p

    # Salutation
    add_body(
        "Dear [Hiring Contact],\n"
        "[Note: You may substitute the personal name of the individual who engaged your services, or retain this respectful salutation.]"
    )

    # Section 1
    add_h1("1. PURPOSE AND CONTEXT OF THIS CORRESPONDENCE")
    add_body(
        "I am writing directly and privately to you regarding our ongoing engagement for the development of the "
        "Chief Red Bear Children's Lodge (CRBCL) software platform. Over the past several months, I have dedicated "
        "substantial effort and specialized engineering to build a custom, enterprise-scale digital platform tailored "
        "to the sovereignty and operational needs of Cowessess First Nation children and families."
    )
    add_body(
        "The purpose of this document is to provide a clear, factual, and professional basis for discussing the financial "
        "arrangements of this project. Specifically, this correspondence seeks to: (1) review the original agreement and quotation; "
        "(2) reconcile the payments received to date; (3) document the extensive expansion of project requirements that occurred "
        "after our initial agreement; (4) address our previous discussions regarding additional compensation; and (5) establish "
        "a constructive, mutually agreed path forward for fair compensation and written clarity regarding remaining work."
    )
    add_body(
        "To protect the professional standing of this project, I have strictly separated technical documentation from commercial "
        "correspondence. The comprehensive technical reports provided for the CRBCL Board of Directors focus exclusively on the platform's "
        "architecture, functionality, and remaining operationalization tasks without any discussion of personal fees or payment arrangements. "
        "This private document is intended exclusively for you as the individual who engaged my services."
    )

    # Section 2
    add_h1("2. THE ORIGINAL AGREEMENT & INITIAL PROJECT QUOTATION")
    add_body(
        "At the beginning of our engagement in early 2026, we discussed an initial software solution to help Chief Red Bear Children's Lodge "
        "transition away from paper binders and basic spreadsheet files. Based on the preliminary and limited project requirements available "
        "at that time, my original quotation was CAD $2,500.00."
    )
    add_body(
        "That initial quotation reflected a modest project scope: a straightforward web-based database application with basic user logins, "
        "a simple client contact index, and basic case note entry. It did not encompass agency-wide multi-departmental workflows, sovereign "
        "First Nations legal modeling, enterprise role-based security, mobile applications, or offline field synchronization."
    )

    # Section 3
    add_h1("3. RECONCILIATION OF PAYMENTS RECEIVED TO DATE")
    add_body(
        "During the course of the initial development and setup, I received progress payments totaling CAD $2,200.00. "
        "I acknowledge and appreciate these disbursements, which supported the foundational stages of the work."
    )
    add_body(
        "Comparing the payments received (CAD $2,200.00) against the original quotation (CAD $2,500.00), there is a nominal "
        "difference of CAD $300.00. I record this figure factually as part of an orderly commercial reconciliation; I do not "
        "characterize this difference as legally overdue without reference to final contract milestone terms, but rather as "
        "the remaining unbilled or unpaid portion of the initial quoted amount."
    )

    # Commercial Facts Table in DOCX
    t_comm = doc.add_table(rows=1 + len(COMMERCIAL_FACTS_TABLE), cols=3)
    t_comm.style = "Table Grid"
    hdr = t_comm.rows[0].cells
    hdr[0].text = "Financial / Commercial Metric"
    hdr[1].text = "Amount / Status"
    hdr[2].text = "Commercial Context & Notes"
    for cell in hdr:
        set_cell_background(cell, PRIMARY_HEX)
        set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for idx, (metric, amt, notes) in enumerate(COMMERCIAL_FACTS_TABLE):
        row = t_comm.rows[idx + 1].cells
        row[0].text = metric
        row[1].text = amt
        row[2].text = notes
        for c_idx, cell in enumerate(row):
            set_cell_background(cell, LIGHT_BG_HEX if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=45, bottom=45, left=60, right=60)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 1:
                p.runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 4
    add_h1("4. SUBSEQUENT SCOPE EXPANSION: ENGINEERING COMPLETED BEYOND ORIGINAL SPECIFICATION")
    add_body(
        "Following our initial agreement, the platform requirements underwent profound structural expansion. As CRBCL leadership "
        "and statutory advisors articulated the full operational mandate of the Cowessess First Nation Miyo-Pimatisiwin Act under "
        "Federal Bill C-92, the project evolved from a simple case-logging utility into a comprehensive, sovereign enterprise platform."
    )
    add_body(
        "I embraced these expanded requirements in good faith, investing hundreds of hours to design, implement, test, and document "
        "capabilities that were entirely outside the original CAD $2,500.00 baseline. A rigorous, evidence-based review of the resulting "
        "codebase confirms the breadth of this additional work:"
    )

    t_exp = doc.add_table(rows=1 + len(EXPANDED_WORK_ITEMS), cols=4)
    t_exp.style = "Table Grid"
    hdr_e = t_exp.rows[0].cells
    hdr_e[0].text = "Functional Area"
    hdr_e[1].text = "Original Baseline"
    hdr_e[2].text = "Engineered Scope Expansion"
    hdr_e[3].text = "Current Technical Status & Evidence"
    for cell in hdr_e:
        set_cell_background(cell, PRIMARY_HEX)
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for idx, (f_area, orig, exp_scope, status_ev) in enumerate(EXPANDED_WORK_ITEMS):
        row = t_exp.rows[idx + 1].cells
        row[0].text = f_area
        row[1].text = orig
        row[2].text = exp_scope
        row[3].text = status_ev
        for c_idx, cell in enumerate(row):
            set_cell_background(cell, LIGHT_BG_HEX if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=40, bottom=40, left=50, right=50)
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(7.5)
            p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 0:
                p.runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_body(
        "Clarification of Implemented vs Unfinished Work: It is vital to maintain complete technical integrity regarding what has "
        "been completed versus what remains pending. The source code, data models, APIs, and automated tests for the above features "
        "are substantially complete. Furthermore, I have personally tested the mobile application on a physical Android device in "
        "offline mode, confirming that offline functionality operates successfully. However, several operationalization items remain "
        "to be completed: (1) live cloud LLM integration for Ask Red Bear (stubbed in source; requires Board vendor selection and API key); "
        "(2) physical vehicle OBD-II GPS telematics hardware procurement; (3) commercial Canadian SMS gateway setup (Twilio); and (4) full "
        "frontline supervisor field acceptance testing for multi-device sync, conflict resolution, and data recovery scenarios."
    )

    # Section 5
    add_h1("5. PREVIOUS DISCUSSIONS REGARDING SCOPE AND ADDITIONAL COMPENSATION")
    add_body(
        "As the development requirements expanded to encompass an enterprise-scale architecture (growing from an initial ~15 tables "
        "into 205 mapped SQLAlchemy models across 14 namespaces and 521 API endpoints), I communicated that the original CAD $2,500.00 "
        "quotation could no longer reasonably cover the expanded platform."
    )
    add_body(
        "During our subsequent discussions, you acknowledged this substantial scope growth and verbally indicated that extra compensation "
        "would be provided to reflect the expanded development. However, no specific revised compensation figure, payment milestone schedule, "
        "or formal written change order was finalized at that time."
    )
    add_body(
        "I state this factually and with complete respect: there is no written dispute or formal claim of breach, but rather a mutual "
        "recognition that the platform built today is an enterprise asset far beyond what either of us contemplated when the original "
        "price was quoted."
    )

    # Section 6
    add_h1("6. CONSTRUCTIVE PATH FORWARD & DISCUSSION FRAMEWORK")
    add_body(
        "I value our working relationship and remain committed to ensuring that Chief Red Bear Children's Lodge receives an exceptional, "
        "operationally successful platform. To bring this engagement to a successful conclusion, I propose that we hold a constructive "
        "discussion focused on the following key points:"
    )
    add_body("• Recognition of Value Delivered: Fair acknowledgment of the extensive engineering completed beyond the initial prototype scope.")
    add_body("• Scope & Status Review: Mutual review of the completed technical capabilities and the specific remaining operationalization tasks.")
    add_body("• Reconciliation of Historical Accounts: Formal reconciliation of the CAD $2,200.00 received and the CAD $300.00 balance.")
    add_body("• Agreement on Fair Additional Compensation: Agreeing upon a reasonable, fair additional compensation amount for the completed out-of-scope work.")
    add_body("• Written Scope for Remaining Work: Setting a defined, written scope and milestones for final UAT support, live deployment, and handover.")
    add_body("• Written Payment Arrangements: Formalizing agreed payment terms and schedules for any further deployment or integration work.")

    # Section 7: Developer Proposal Placeholder
    add_h1("7. PROPOSAL FOR ADDITIONAL COMPENSATION (DEVELOPER INSERTION SECTION)")
    add_body(
        "To facilitate our discussion, I have set aside this section for outlining a proposed compensation structure. "
        "I invite us to review these parameters and agree on an arrangement that is fair to both parties:"
    )

    prop_tbl = doc.add_table(rows=1, cols=1)
    prop_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    pc = prop_tbl.cell(0, 0)
    set_cell_background(pc, "F1F5F9")
    set_cell_margins(pc, top=140, bottom=140, left=180, right=180)
    pp = pc.paragraphs[0]
    pr1 = pp.add_run("========================================================================================\n")
    pr1.bold = True
    pr1.font.name = "Courier"
    pr1.font.size = Pt(8.5)
    pr1.font.color.rgb = RGBColor(100, 116, 139)

    pr2 = pp.add_run("[PROPOSED ADDITIONAL COMPENSATION & PAYMENT TERMS — DEVELOPER COMPLETION SECTION]\n\n")
    pr2.bold = True
    pr2.font.name = "Arial"
    pr2.font.size = Pt(10)
    pr2.font.color.rgb = RGBColor(107, 29, 47)

    pr3 = pp.add_run(
        "1. Proposed Additional Compensation for Completed Out-of-Scope Work:\n"
        "   CAD $ [DEVELOPER TO INSERT PROPOSED FIGURE, E.G., CAD $X,XXX.XX]\n\n"
        "2. Proposed Basis / Payment Structure:\n"
        "   [e.g., Lump-sum adjustment, phased milestone payments, or agreed fixed fee for Phase 1-5 deliverables]\n\n"
        "3. Proposed Terms for Remaining Go-Live Operationalization:\n"
        "   [e.g., $X upon successful completion of UAT and live deployment; $Y upon staff training support]\n\n"
        "4. Payment Schedule & Milestones:\n"
        "   [INSERT PROPOSED DATES OR MILESTONE TRIGGERS]\n\n"
        "5. Additional Notes & Considerations:\n"
        "   [INSERT ANY SPECIFIC WORKING ARRANGEMENTS OR TIMELINES]\n"
    )
    pr3.font.name = "Arial"
    pr3.font.size = Pt(9)
    pr3.font.color.rgb = RGBColor(30, 41, 59)

    pr4 = pp.add_run("========================================================================================")
    pr4.bold = True
    pr4.font.name = "Courier"
    pr4.font.size = Pt(8.5)
    pr4.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 8
    add_h1("8. PROFESSIONAL COMMITMENT & CONCLUSION")
    add_body(
        "I have invested a tremendous amount of care and technical dedication into this project because I believe deeply "
        "in the mission of Chief Red Bear Children's Lodge and the importance of Indigenous-governed child and family wellness. "
        "All source code, database migrations, configuration files, and technical documentation are thoroughly organized, versioned, "
        "and preserved in the project repository."
    )
    add_body(
        "I am ready to engage in an open, constructive, and amicable discussion at your earliest convenience to review this proposal "
        "and establish a mutually beneficial agreement. Thank you for your consideration and for your partnership in bringing this "
        "platform to its current state of technical maturity."
    )
    add_body("\nSincerely,\n\nLead Software Developer\nChief Red Bear Children's Lodge Platform Project")

    doc.save(docx_path)
    print(f"Private Commercial DOCX written to: {docx_path}")

    # =========================================================================
    # BUILD PDF
    # =========================================================================
    styles = get_pdf_styles_v2()
    pdf_doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    story = []

    # Confidential Notice Box
    notice_text = (
        "<b>PRIVATE AND CONFIDENTIAL — FOR THE ORIGINAL HIRING CONTACT ONLY</b><br/><br/>"
        "This correspondence is strictly private and commercial in nature. It is addressed exclusively to the "
        "individual who engaged the developer's services for the Chief Red Bear Children's Lodge platform. "
        "It must not be circulated to the CRBCL Board of Directors, executive committees, or external evaluators. "
        "Technical and progress documentation for the Board has been provided separately in dedicated Board reports."
    )
    from scripts.deliverables_generator_v2.common_v2 import create_pdf_notice_box_v2
    story.append(create_pdf_notice_box_v2("CONFIDENTIAL COMMERCIAL CORRESPONDENCE", notice_text, styles))
    story.append(Spacer(1, 14))

    # Header Titles
    story.append(Paragraph("CHIEF RED BEAR CHILDREN'S LODGE (CRBCL) PLATFORM<br/>COMMERCIAL REVIEW & COMPENSATION DISCUSSION", styles["CRBCL_CoverTitle"]))
    story.append(Paragraph("Original Agreement, Scope Expansion, Completed Engineering, and Outstanding Compensation Discussion", styles["CRBCL_CoverSubtitle"]))
    story.append(Paragraph("<b>Date:</b> October 2026 &nbsp;|&nbsp; <b>Context:</b> Private Developer Commercial Correspondence &nbsp;|&nbsp; <b>Audience:</b> Original Hiring Contact Only", styles["CRBCL_CoverMeta"]))
    story.append(HRFlowable(width="100%", thickness=1.5, color=COLOR_PRIMARY, spaceAfter=14))

    # Salutation
    story.append(Paragraph("<b>Dear [Hiring Contact],</b><br/><i>[Note: You may substitute the personal name of the individual who engaged your services, or retain this respectful salutation.]</i>", styles["CRBCL_Body"]))
    story.append(Spacer(1, 8))

    # Section 1
    story.append(Paragraph("1. PURPOSE AND CONTEXT OF THIS CORRESPONDENCE", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "I am writing directly and privately to you regarding our ongoing engagement for the development of the "
        "Chief Red Bear Children's Lodge (CRBCL) software platform. Over the past several months, I have dedicated "
        "substantial effort and specialized engineering to build a custom, enterprise-scale digital platform tailored "
        "to the sovereignty and operational needs of Cowessess First Nation children and families.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "The purpose of this document is to provide a clear, factual, and professional basis for discussing the financial "
        "arrangements of this project. Specifically, this correspondence seeks to: (1) review the original agreement and quotation; "
        "(2) reconcile the payments received to date; (3) document the extensive expansion of project requirements that occurred "
        "after our initial agreement; (4) address our previous discussions regarding additional compensation; and (5) establish "
        "a constructive, mutually agreed path forward for fair compensation and written clarity regarding remaining work.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "To protect the professional standing of this project, I have strictly separated technical documentation from commercial "
        "correspondence. The comprehensive technical reports provided for the CRBCL Board of Directors focus exclusively on the platform's "
        "architecture, functionality, and remaining operationalization tasks without any discussion of personal fees or payment arrangements. "
        "This private document is intended exclusively for you as the individual who engaged my services.",
        styles["CRBCL_Body"]
    ))
    story.append(Spacer(1, 8))

    # Section 2
    story.append(Paragraph("2. THE ORIGINAL AGREEMENT & INITIAL PROJECT QUOTATION", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "At the beginning of our engagement in early 2026, we discussed an initial software solution to help Chief Red Bear Children's Lodge "
        "transition away from paper binders and basic spreadsheet files. Based on the preliminary and limited project requirements available "
        "at that time, my original quotation was <b>CAD $2,500.00</b>.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "That initial quotation reflected a modest project scope: a straightforward web-based database application with basic user logins, "
        "a simple client contact index, and basic case note entry. It did not encompass agency-wide multi-departmental workflows, sovereign "
        "First Nations legal modeling, enterprise role-based security, mobile applications, or offline field synchronization.",
        styles["CRBCL_Body"]
    ))
    story.append(Spacer(1, 8))

    # Section 3
    story.append(Paragraph("3. RECONCILIATION OF PAYMENTS RECEIVED TO DATE", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "During the course of the initial development and setup, I received progress payments totaling <b>CAD $2,200.00</b>. "
        "I acknowledge and appreciate these disbursements, which supported the foundational stages of the work.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "Comparing the payments received (CAD $2,200.00) against the original quotation (CAD $2,500.00), there is a nominal "
        "difference of <b>CAD $300.00</b>. I record this figure factually as part of an orderly commercial reconciliation; I do not "
        "characterize this difference as legally overdue without reference to final contract milestone terms, but rather as "
        "the remaining unbilled or unpaid portion of the initial quoted amount.",
        styles["CRBCL_Body"]
    ))

    # PDF Commercial Facts Table
    pdf_comm_rows = [
        [
            Paragraph("<b>Financial / Commercial Metric</b>", styles["CRBCL_TableHeader"]),
            Paragraph("<b>Amount / Status</b>", styles["CRBCL_TableHeader"]),
            Paragraph("<b>Commercial Context & Notes</b>", styles["CRBCL_TableHeader"]),
        ]
    ]
    for m, a, n in COMMERCIAL_FACTS_TABLE:
        pdf_comm_rows.append([
            Paragraph(f"<b>{m}</b>", styles["CRBCL_TableCellBold"]),
            Paragraph(f"<b>{a}</b>", styles["CRBCL_TableCellBold"]),
            Paragraph(n, styles["CRBCL_TableCell"]),
        ])
    t_comm_pdf = Table(pdf_comm_rows, colWidths=[150, 110, 244])
    t_comm_pdf.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_PRIMARY),
        ("GRID", (0, 0), (-1, -1), 0.5, COLOR_BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [COLOR_LIGHT_BG, colors.white]),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_comm_pdf)
    story.append(Spacer(1, 10))
    story.append(PageBreak())

    # Section 4
    story.append(Paragraph("4. SUBSEQUENT SCOPE EXPANSION: ENGINEERING COMPLETED BEYOND ORIGINAL SPECIFICATION", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "Following our initial agreement, the platform requirements underwent profound structural expansion. As CRBCL leadership "
        "and statutory advisors articulated the full operational mandate of the Cowessess First Nation Miyo-Pimatisiwin Act under "
        "Federal Bill C-92, the project evolved from a simple case-logging utility into a comprehensive, sovereign enterprise platform.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "I embraced these expanded requirements in good faith, investing hundreds of hours to design, implement, test, and document "
        "capabilities that were entirely outside the original CAD $2,500.00 baseline. A rigorous, evidence-based review of the resulting "
        "codebase confirms the breadth of this additional work:",
        styles["CRBCL_Body"]
    ))

    # PDF Expanded Work Items Table
    pdf_exp_rows = [
        [
            Paragraph("<b>Functional Area</b>", styles["CRBCL_TableHeader"]),
            Paragraph("<b>Original Baseline</b>", styles["CRBCL_TableHeader"]),
            Paragraph("<b>Engineered Scope Expansion</b>", styles["CRBCL_TableHeader"]),
            Paragraph("<b>Current Technical Status & Evidence</b>", styles["CRBCL_TableHeader"]),
        ]
    ]
    for fa, ob, es, ce in EXPANDED_WORK_ITEMS:
        pdf_exp_rows.append([
            Paragraph(f"<b>{fa}</b>", styles["CRBCL_TableCellBold"]),
            Paragraph(ob, styles["CRBCL_TableCell"]),
            Paragraph(es, styles["CRBCL_TableCell"]),
            Paragraph(ce, styles["CRBCL_TableCell"]),
        ])
    t_exp_pdf = Table(pdf_exp_rows, colWidths=[110, 100, 160, 134])
    t_exp_pdf.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_PRIMARY),
        ("GRID", (0, 0), (-1, -1), 0.5, COLOR_BORDER),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [COLOR_LIGHT_BG, colors.white]),
        ("PADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(t_exp_pdf)
    story.append(Spacer(1, 8))

    story.append(Paragraph(
        "<b>Clarification of Implemented vs Unfinished Work:</b> It is vital to maintain complete technical integrity regarding what has "
        "been completed versus what remains pending. The source code, data models, APIs, and automated tests for the above features "
        "are substantially complete. Furthermore, I have personally tested the mobile application on a physical Android device in "
        "offline mode, confirming that offline functionality operates successfully. However, several operationalization items remain "
        "to be completed: (1) live cloud LLM integration for Ask Red Bear (stubbed in source; requires Board vendor selection and API key); "
        "(2) physical vehicle OBD-II GPS telematics hardware procurement; (3) commercial Canadian SMS gateway setup (Twilio); and (4) full "
        "frontline supervisor field acceptance testing for multi-device sync, conflict resolution, and data recovery scenarios.",
        styles["CRBCL_Body"]
    ))
    story.append(Spacer(1, 10))
    story.append(PageBreak())

    # Section 5
    story.append(Paragraph("5. PREVIOUS DISCUSSIONS REGARDING SCOPE AND ADDITIONAL COMPENSATION", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "As the development requirements expanded to encompass an enterprise-scale architecture (growing from an initial ~15 tables "
        "into 205 mapped SQLAlchemy models across 14 namespaces and 521 API endpoints), I communicated that the original CAD $2,500.00 "
        "quotation could no longer reasonably cover the expanded platform.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "During our subsequent discussions, you acknowledged this substantial scope growth and verbally indicated that extra compensation "
        "would be provided to reflect the expanded development. However, no specific revised compensation figure, payment milestone schedule, "
        "or formal written change order was finalized at that time.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "I state this factually and with complete respect: there is no written dispute or formal claim of breach, but rather a mutual "
        "recognition that the platform built today is an enterprise asset far beyond what either of us contemplated when the original "
        "price was quoted.",
        styles["CRBCL_Body"]
    ))
    story.append(Spacer(1, 8))

    # Section 6
    story.append(Paragraph("6. CONSTRUCTIVE PATH FORWARD & DISCUSSION FRAMEWORK", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "I value our working relationship and remain committed to ensuring that Chief Red Bear Children's Lodge receives an exceptional, "
        "operationally successful platform. To bring this engagement to a successful conclusion, I propose that we hold a constructive "
        "discussion focused on the following key points:",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph("• <b>Recognition of Value Delivered:</b> Fair acknowledgment of the extensive engineering completed beyond the initial prototype scope.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Scope & Status Review:</b> Mutual review of the completed technical capabilities and the specific remaining operationalization tasks.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Reconciliation of Historical Accounts:</b> Formal reconciliation of the CAD $2,200.00 received and the CAD $300.00 balance.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Agreement on Fair Additional Compensation:</b> Agreeing upon a reasonable, fair additional compensation amount for the completed out-of-scope work.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Written Scope for Remaining Work:</b> Setting a defined, written scope and milestones for final UAT support, live deployment, and handover.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Written Payment Arrangements:</b> Formalizing agreed payment terms and schedules for any further deployment or integration work.", styles["CRBCL_Body"]))
    story.append(Spacer(1, 8))

    # Section 7: Developer Proposal Placeholder
    story.append(Paragraph("7. PROPOSAL FOR ADDITIONAL COMPENSATION (DEVELOPER INSERTION SECTION)", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "To facilitate our discussion, I have set aside this section for outlining a proposed compensation structure. "
        "I invite us to review these parameters and agree on an arrangement that is fair to both parties:",
        styles["CRBCL_Body"]
    ))

    prop_box_text = (
        "<b>========================================================================================</b><br/>"
        "<b>[PROPOSED ADDITIONAL COMPENSATION & PAYMENT TERMS — DEVELOPER COMPLETION SECTION]</b><br/><br/>"
        "<b>1. Proposed Additional Compensation for Completed Out-of-Scope Work:</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;CAD $ [DEVELOPER TO INSERT PROPOSED FIGURE, E.G., CAD $X,XXX.XX]<br/><br/>"
        "<b>2. Proposed Basis / Payment Structure:</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;[e.g., Lump-sum adjustment, phased milestone payments, or agreed fixed fee for Phase 1-5 deliverables]<br/><br/>"
        "<b>3. Proposed Terms for Remaining Go-Live Operationalization:</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;[e.g., $X upon successful completion of UAT and live deployment; $Y upon staff training support]<br/><br/>"
        "<b>4. Payment Schedule & Milestones:</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;[INSERT PROPOSED DATES OR MILESTONE TRIGGERS]<br/><br/>"
        "<b>5. Additional Notes & Considerations:</b><br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;[INSERT ANY SPECIFIC WORKING ARRANGEMENTS OR TIMELINES]<br/>"
        "<b>========================================================================================</b>"
    )
    t_prop_pdf = Table([[Paragraph(prop_box_text, styles["CRBCL_Body"])]], colWidths=[504])
    t_prop_pdf.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
        ("BOX", (0, 0), (-1, -1), 1, COLOR_PRIMARY),
        ("PADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(t_prop_pdf)
    story.append(Spacer(1, 10))

    # Section 8
    story.append(Paragraph("8. PROFESSIONAL COMMITMENT & CONCLUSION", styles["CRBCL_H1"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COLOR_PRIMARY, spaceAfter=6))
    story.append(Paragraph(
        "I have invested a tremendous amount of care and technical dedication into this project because I believe deeply "
        "in the mission of Chief Red Bear Children's Lodge and the importance of Indigenous-governed child and family wellness. "
        "All source code, database migrations, configuration files, and technical documentation are thoroughly organized, versioned, "
        "and preserved in the project repository.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph(
        "I am ready to engage in an open, constructive, and amicable discussion at your earliest convenience to review this proposal "
        "and establish a mutually beneficial agreement. Thank you for your consideration and for your partnership in bringing this "
        "platform to its current state of technical maturity.",
        styles["CRBCL_Body"]
    ))
    story.append(Spacer(1, 10))
    story.append(Paragraph("<b>Sincerely,</b><br/><br/>Lead Software Developer<br/>Chief Red Bear Children's Lodge Platform Project", styles["CRBCL_Body"]))

    pdf_doc.build(story, canvasmaker=NumberedCanvasPrivate)
    print(f"Private Commercial PDF written to: {pdf_path}")
    return docx_path, pdf_path


if __name__ == "__main__":
    build_private_commercial_doc()
