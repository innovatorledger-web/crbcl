"""Generate Document 2: Updated Technical Architecture and Development Documentation (DOCX & PDF)."""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib import colors

from scripts.deliverables_generator.common import (
    DELIVERABLES_DIR, NumberedCanvas, get_pdf_styles, create_pdf_notice_box,
    PRIMARY_HEX, SECONDARY_HEX, DARK_HEX, LIGHT_BG_HEX, BORDER_HEX,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_DARK, COLOR_LIGHT_BG, COLOR_BORDER,
    set_cell_background, set_cell_margins
)

DOC_NAME_BASE = "CRBCL_Updated_Technical_Architecture_and_Development_Documentation_October_2026"


def build_docx():
    docx_path = os.path.join(DELIVERABLES_DIR, f"{DOC_NAME_BASE}.docx")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = "CHIEF RED BEAR CHILDREN'S LODGE — TECHNICAL ARCHITECTURE DOCUMENTATION (OCTOBER 2026)"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.style.font.name = "Arial"
    hp.style.font.size = Pt(8.5)
    hp.style.font.color.rgb = RGBColor(100, 116, 139)

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = "CONFIDENTIAL — TECHNICAL SPECIFICATION & HANDOVER REFERENCE"
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.style.font.name = "Arial"
    fp.style.font.size = Pt(8.5)
    fp.style.font.color.rgb = RGBColor(100, 116, 139)

    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("CRBCL PLATFORM TECHNICAL ARCHITECTURE\n& SYSTEM SPECIFICATION")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(19)
    r_title.font.color.rgb = RGBColor(107, 29, 47)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Comprehensive Architectural Inventory, Domain Schema, Integration Models, and Technical Handover Reference")
    r_sub.bold = True
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(201, 138, 44)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("Cowessess First Nation | Chief Red Bear Children's Lodge\nPreparation Date: October 27, 2026 | Commit HEAD: 0746579 | Branch: main\nArchitecture Baseline: Post-Phase 5 Release Audit")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_sec(h1_text):
        p = doc.add_paragraph()
        r = p.add_run(h1_text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(13.5)
        r.font.color.rgb = RGBColor(107, 29, 47)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)

    def add_h2(h2_text):
        p = doc.add_paragraph()
        r = p.add_run(h2_text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(9.5)
            rb.font.color.rgb = RGBColor(30, 41, 59)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_after = Pt(5)

    # Section 1: System Topology
    add_sec("1. SYSTEM TOPOLOGY & DECOUPLED ARCHITECTURE")
    add_body(
        "The Chief Red Bear Children's Lodge (CRBCL) Family Wellness Case Management Platform is engineered as a modern, decoupled, multi-tier software system. The architecture separates presentation, API orchestration, persistent relational storage, and background processing to guarantee high data availability, sovereign access controls, and horizontal scalability."
    )
    add_body(
        "1. Web Presentation Tier: Built with React 18, Vite, and Tailwind CSS. The web client communicates with the backend via a centralized Axios/Fetch API client implementing automated JWT token refresh, unified error envelopes, and client-side route guards.",
        "• "
    )
    add_body(
        "2. Mobile Presentation Tier: Cross-platform client built in Flutter / Dart. Provides offline-first capabilities for rural caseworkers using local SQLite caching, local biometric authentication, and an outbox queue engine for bidirectional sync.",
        "• "
    )
    add_body(
        "3. API & Orchestration Tier: High-performance asynchronous Python service built on FastAPI and Pydantic v2. Implements dependency-injected database sessions, fine-grained RBAC permission enforcement, and an Authorization-First AI context gateway.",
        "• "
    )
    add_body(
        "4. Relational Persistence Tier: PostgreSQL 15+ managed instance utilizing SQLAlchemy 2.0 async engine with connection pooling. Database migrations are rigorously version-controlled through 31 sequential Alembic revision scripts.",
        "• "
    )

    # Section 2: Architectural Evolution
    add_sec("2. ARCHITECTURAL EVOLUTION: ORIGINAL SPECIFICATION VS CURRENT IMPLEMENTATION")
    add_body(
        "A rigorous review of the original CRBCL Master Project Document reveals significant positive architectural evolution. Rather than degrading the design, implementation changes were made to enhance First Nations data sovereignty, security, and offline field utility:"
    )

    evol_table = doc.add_table(rows=6, cols=3)
    evol_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["System Domain", "Original Specification Baseline", "Current Production Architecture (Commit 0746579)"]
    for i, h in enumerate(headers):
        cell = evol_table.cell(0, i)
        set_cell_background(cell, PRIMARY_HEX)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    evol_data = [
        ("Database & Schema", "Monolithic schema with shared client-case tables", "Decoupled domain models across 14 distinct namespaces with versioned audit histories (126 models)."),
        ("API Architecture", "Synchronous Flask/Django endpoints", "Asynchronous FastAPI with non-blocking asyncpg connections, Pydantic v2 validation, and OpenAPI 3.1 schemas."),
        ("AI / Ask Red Bear", "Generic chat interface connected to cloud LLM", "Strict Authorization-First AI Gateway, prompt guardrails, field-level redactions, and simulated stub engine pending CRBCL approval."),
        ("Offline Field Sync", "Standard web responsive views only", "Dedicated native Flutter client with SQLite offline drafting, encrypted local storage, and transactional sync engine."),
        ("Data Import & Migration", "Direct bulk SQL insert from spreadsheets", "Isolated quarantine staging wizard, 3-tier duplicate protection engine, client reconciliation review queue, and inactive staff protection."),
    ]

    for row_idx, row_data in enumerate(evol_data, start=1):
        bg = LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = evol_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Data Layer
    add_sec("3. RELATIONAL DATA LAYER & DOMAIN SCHEMA ARCHITECTURE")
    add_body(
        "The relational schema encompasses 126 SQLAlchemy models distributed across 14 business domains. All primary keys utilize RFC 4122 Version 4 UUIDs to prevent ID enumeration attacks and facilitate distributed offline generation. Foreign keys enforce referential integrity with explicit ondelete cascading or restrict rules. Below is the domain mapping breakdown:"
    )

    schema_table = doc.add_table(rows=8, cols=3)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Domain Namespace", "Primary Tables & Entities", "Relational Integrity & Audit Controls"]
    for i, h in enumerate(s_headers):
        cell = schema_table.cell(0, i)
        set_cell_background(cell, PRIMARY_HEX)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    s_data = [
        ("Identity & Access (RBAC)", "users, roles, permissions, role_permissions, user_roles, audit_log", "Granular 114 permissions; immutable audit trails on all permission grants."),
        ("Person & Client Master", "persons, clients, families, person_identifiers, client_approvals", "Fuzzy duplicate detection tables; controlled merge history; family linkage graph."),
        ("Case Management", "cases, case_assignments, case_restrictions, case_links, case_status_history", "Mandatory active assignment check; case restriction suppression rules."),
        ("Clinical Notes & Evidence", "case_notes, case_note_addenda, case_note_attachments, note_persons", "Immutable note locking after sign-off; addendum versioning; no hard deletes."),
        ("Assessments & Safety", "assessments, assessment_sections, questions, answers, safety_plans", "Multi-version template engine; supervisor unlock tracking; sequence numbering."),
        ("Prevention & Healing", "ph_records, ph_intake_links, ph_staff_quarantine, ph_duplicate_queue", "3-tier duplicate scoring; provenance lineage; isolated quarantine tables."),
        ("Operations, Fleet & Finance", "invoices, service_requests, billing_rates, vehicles, fleet_trips, geofences", "Multi-fund ledger; strict invoice immutability; vehicle telematics tracking."),
    ]

    for row_idx, sd in enumerate(s_data, start=1):
        bg = LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(sd):
            cell = schema_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 4: Security Architecture
    add_sec("4. AUTHENTICATION, RBAC & FIRST NATIONS DATA SOVEREIGNTY")
    add_body(
        "Security is implemented as a multi-layered defense-in-depth model aligned with Indigenous OCAP® (Ownership, Control, Access, Possession) principles and ISO 27001 standards:"
    )
    add_body(
        "1. Dual Authentication Credentials: The system supports both standard email/password authentication (Argon2id hashing) and frontline 6-digit secure PIN authentication. PIN authentication enables rapid, touch-friendly access for caseworkers on mobile devices or shared terminals while maintaining individual session auditing.",
        "• "
    )
    add_body(
        "2. Granular 114-Permission RBAC Matrix: Rather than coarse role checking (e.g. 'is_admin'), every single endpoint and database action validates discrete permissions (e.g., 'case:create', 'client:medical:read', 'finance:invoice:approve'). Fourteen canonical roles provide pre-packaged role bundles.",
        "• "
    )
    add_body(
        "3. Case-Level Conflict Restrictions: In Indigenous community contexts, caseworkers may have personal or familial relationships with clients. The platform implements an active Case Restriction model: any caseworker assigned a restriction on a specific case is completely blinded from seeing the case in search results, dashboards, audit reports, or Ask Red Bear AI queries.",
        "• "
    )
    add_body(
        "4. Staff Login Quarantine: Historical staff records extracted from legacy spreadsheets are imported into an inactive quarantine status. The application actively prevents password/PIN generation or session issuance for quarantined records until the System Administrator formally verifies the employee.",
        "• "
    )

    # Section 5: Ask Red Bear Deep Dive
    add_sec("5. ASK RED BEAR AI ARCHITECTURE DEEP-DIVE")
    add_body(
        "Ask Red Bear is engineered around the Gateway and Context Manager patterns to guarantee that AI assistance never bypasses data confidentiality or statutory child welfare rules:"
    )
    add_body(
        "Prompt Guard & Safety Filter: Inbound caseworker prompts are scanned for prompt injection attacks, role breakout attempts, or illegal data extraction requests before passing to context evaluation.",
        "Step 1 — "
    )
    add_body(
        "Authorization-First Context Construction: The Gateway queries the user's active permissions and active case restrictions. Only records to which the user has explicit, un-restricted access are pulled into the retrieval context.",
        "Step 2 — "
    )
    add_body(
        "Allowlisted Tool Execution: If the prompt requests a structured query (e.g., 'Summarize my active cases'), the Gateway invokes internal deterministic tools (e.g., get_case_summary, get_overdue_goals) rather than allowing the LLM to hallucinate database contents.",
        "Step 3 — "
    )
    add_body(
        "Field-Level Redaction Engine: Sensitive attributes (child medical notes, allergies, reporter names/phone numbers, financial ledger details) are programmatically stripped from context documents unless the user possesses explicit elevated permissions.",
        "Step 4 — "
    )
    add_body(
        "Speech Transcription as Input Aid: Voice audio recorded via browser or mobile is transcribed in memory and returned to the text box for user editing. Audio is never automatically committed to client records.",
        "Step 5 — "
    )
    add_body(
        "Current LLM Engine State: Currently uses FakeAiProvider in the backend. Transitioning to live production requires provisioning an enterprise API key for a Canadian-hosted sovereign LLM (e.g., Claude 3.5 Sonnet on AWS Canada Central).",
        "Step 6 — "
    )

    # Section 6: Mobile & Offline
    add_sec("6. MOBILE FIELD CLIENT & OFFLINE SYNCHRONIZATION")
    add_body(
        "The Flutter mobile app enables caseworkers to work in remote reserve areas where cell coverage is nonexistent. When disconnected, workers can view cached active cases, draft new case notes, record voice memos, and conduct home visit assessments. All offline mutations are written to an encrypted local SQLite outbox table with idempotency keys. Upon reconnecting to WiFi/cellular, the SyncEngine processes the queue sequentially, resolving version conflicts using server-wins timestamps and notifying the user of successful commits."
    )

    # Section 7: Technical Debt & Next Steps
    add_sec("7. TECHNICAL DEBT, SYSTEM RISKS & HANDOVER READINESS")
    add_body(
        "For an independent software engineering team evaluating this codebase, the technical debt is low and architecture adheres to modern best practices. The remaining tasks are operationalization items:"
    )
    add_body("1. Staging DNS & CORS Synchronization: Ensure Railway backend production headers seamlessly align with the Vercel/DigitalOcean production frontend domain (https://genserver.online).", "• ")
    add_body("2. Third-Party Gateway Credential Setup: Provide live production credentials for Twilio Canada (SMS notifications), Geotab (telematics), and Anthropic/Azure (AI LLM).", "• ")
    add_body("3. End-to-End Browser UAT Automation: Execute Cypress or Playwright test suites against the live domain to establish baseline browser interaction traces.", "• ")

    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    doc.save(docx_path)
    print(f"Document 2 DOCX successfully written to: {docx_path}")
    return docx_path


def build_pdf():
    pdf_path = os.path.join(DELIVERABLES_DIR, f"{DOC_NAME_BASE}.pdf")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = get_pdf_styles()
    story = []

    story.append(Paragraph("CRBCL PLATFORM TECHNICAL ARCHITECTURE<br/>& SYSTEM SPECIFICATION", styles["CRBCL_DocTitle"]))
    story.append(Paragraph("Comprehensive Architectural Inventory, Domain Schema, Integration Models, and Technical Handover Reference", styles["CRBCL_DocSubtitle"]))
    story.append(Paragraph("Cowessess First Nation | Chief Red Bear Children's Lodge<br/>Preparation Date: October 27, 2026 | Commit HEAD: 0746579 | Branch: main<br/>Architecture Baseline: Post-Phase 5 Release Audit", styles["CRBCL_MetaText"]))
    story.append(HRFlowable(width="100%", thickness=1, color=COLOR_PRIMARY, spaceAfter=14))

    # Section 1
    story.append(Paragraph("1. SYSTEM TOPOLOGY & DECOUPLED ARCHITECTURE", styles["CRBCL_H1"]))
    story.append(Paragraph(
        "The Chief Red Bear Children's Lodge (CRBCL) Family Wellness Case Management Platform is engineered as a modern, decoupled, multi-tier software system. The architecture separates presentation, API orchestration, persistent relational storage, and background processing to guarantee high data availability, sovereign access controls, and horizontal scalability.",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph("• <b>Web Presentation Tier:</b> Built with React 18, Vite, and Tailwind CSS. The web client communicates with the backend via a centralized API client implementing automated JWT refresh, unified error envelopes, and client-side route guards.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Mobile Presentation Tier:</b> Cross-platform client built in Flutter / Dart. Provides offline-first capabilities for rural caseworkers using local SQLite caching, local biometric authentication, and an outbox queue engine for bidirectional sync.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>API & Orchestration Tier:</b> High-performance asynchronous Python service built on FastAPI and Pydantic v2. Implements dependency-injected database sessions, fine-grained RBAC permission enforcement, and an Authorization-First AI context gateway.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Relational Persistence Tier:</b> PostgreSQL 15+ managed instance utilizing SQLAlchemy 2.0 async engine with connection pooling. Database migrations are rigorously version-controlled through 31 sequential Alembic revision scripts.", styles["CRBCL_Body"]))
    story.append(Spacer(1, 10))

    # Section 2
    story.append(Paragraph("2. ARCHITECTURAL EVOLUTION: ORIGINAL SPECIFICATION VS CURRENT IMPLEMENTATION", styles["CRBCL_H1"]))
    story.append(Paragraph(
        "A rigorous review of the original CRBCL Master Project Document reveals significant positive architectural evolution. Rather than degrading the design, implementation changes were made to enhance First Nations data sovereignty, security, and offline field utility:",
        styles["CRBCL_Body"]
    ))

    evol_rows = [
        [Paragraph("<b>System Domain</b>", styles["CRBCL_TableHeader"]), Paragraph("<b>Original Specification Baseline</b>", styles["CRBCL_TableHeader"]), Paragraph("<b>Current Production Architecture (0746579)</b>", styles["CRBCL_TableHeader"])],
        [Paragraph("<b>Database & Schema</b>", styles["CRBCL_TableCellBold"]), Paragraph("Monolithic schema with shared client-case tables", styles["CRBCL_TableCell"]), Paragraph("Decoupled domain models across 14 distinct namespaces with versioned audit histories (126 models).", styles["CRBCL_TableCell"])],
        [Paragraph("<b>API Architecture</b>", styles["CRBCL_TableCellBold"]), Paragraph("Synchronous Flask/Django endpoints", styles["CRBCL_TableCell"]), Paragraph("Asynchronous FastAPI with non-blocking asyncpg connections, Pydantic v2 validation, and OpenAPI 3.1 schemas.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>AI / Ask Red Bear</b>", styles["CRBCL_TableCellBold"]), Paragraph("Generic chat interface connected to cloud LLM", styles["CRBCL_TableCell"]), Paragraph("Strict Authorization-First AI Gateway, prompt guardrails, field redactions, and simulated stub engine.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Offline Field Sync</b>", styles["CRBCL_TableCellBold"]), Paragraph("Standard web responsive views only", styles["CRBCL_TableCell"]), Paragraph("Dedicated native Flutter client with SQLite offline drafting, encrypted storage, and transactional sync engine.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Data Import & Migration</b>", styles["CRBCL_TableCellBold"]), Paragraph("Direct bulk SQL insert from spreadsheets", styles["CRBCL_TableCell"]), Paragraph("Isolated quarantine staging wizard, 3-tier duplicate engine, client reconciliation review queue, inactive staff protection.", styles["CRBCL_TableCell"])],
    ]

    t_evol = Table(evol_rows, colWidths=[110, 180, 214])
    t_evol.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_PRIMARY),
        ("GRID", (0, 0), (-1, -1), 0.5, COLOR_BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [COLOR_LIGHT_BG, colors.white]),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t_evol)
    story.append(Spacer(1, 14))

    # Section 3
    story.append(Paragraph("3. RELATIONAL DATA LAYER & DOMAIN SCHEMA ARCHITECTURE", styles["CRBCL_H1"]))
    story.append(Paragraph(
        "The relational schema encompasses 126 SQLAlchemy models distributed across 14 business domains. All primary keys utilize RFC 4122 Version 4 UUIDs to prevent ID enumeration attacks and facilitate distributed offline generation. Foreign keys enforce referential integrity with explicit ondelete cascading or restrict rules:",
        styles["CRBCL_Body"]
    ))

    schema_rows = [
        [Paragraph("<b>Domain Namespace</b>", styles["CRBCL_TableHeader"]), Paragraph("<b>Primary Tables & Entities</b>", styles["CRBCL_TableHeader"]), Paragraph("<b>Relational Integrity & Audit Controls</b>", styles["CRBCL_TableHeader"])],
        [Paragraph("<b>Identity & Access (RBAC)</b>", styles["CRBCL_TableCellBold"]), Paragraph("users, roles, permissions, role_permissions, user_roles, audit_log", styles["CRBCL_TableCell"]), Paragraph("Granular 114 permissions; immutable audit trails on all permission grants.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Person & Client Master</b>", styles["CRBCL_TableCellBold"]), Paragraph("persons, clients, families, person_identifiers, client_approvals", styles["CRBCL_TableCell"]), Paragraph("Fuzzy duplicate detection tables; controlled merge history; family linkage graph.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Case Management</b>", styles["CRBCL_TableCellBold"]), Paragraph("cases, case_assignments, case_restrictions, case_links, case_status_history", styles["CRBCL_TableCell"]), Paragraph("Mandatory active assignment check; case restriction suppression rules.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Clinical Notes & Evidence</b>", styles["CRBCL_TableCellBold"]), Paragraph("case_notes, case_note_addenda, case_note_attachments, note_persons", styles["CRBCL_TableCell"]), Paragraph("Immutable note locking after sign-off; addendum versioning; no hard deletes.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Assessments & Safety</b>", styles["CRBCL_TableCellBold"]), Paragraph("assessments, assessment_sections, questions, answers, safety_plans", styles["CRBCL_TableCell"]), Paragraph("Multi-version template engine; supervisor unlock tracking; sequence numbering.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Prevention & Healing</b>", styles["CRBCL_TableCellBold"]), Paragraph("ph_records, ph_intake_links, ph_staff_quarantine, ph_duplicate_queue", styles["CRBCL_TableCell"]), Paragraph("3-tier duplicate scoring; provenance lineage; isolated quarantine tables.", styles["CRBCL_TableCell"])],
        [Paragraph("<b>Operations, Fleet & Finance</b>", styles["CRBCL_TableCellBold"]), Paragraph("invoices, service_requests, billing_rates, vehicles, fleet_trips, geofences", styles["CRBCL_TableCell"]), Paragraph("Multi-fund ledger; strict invoice immutability; vehicle telematics tracking.", styles["CRBCL_TableCell"])],
    ]

    t_schema = Table(schema_rows, colWidths=[120, 190, 194])
    t_schema.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_PRIMARY),
        ("GRID", (0, 0), (-1, -1), 0.5, COLOR_BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [COLOR_LIGHT_BG, colors.white]),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_schema)
    story.append(Spacer(1, 14))

    # Section 4
    story.append(Paragraph("4. AUTHENTICATION, RBAC & FIRST NATIONS DATA SOVEREIGNTY", styles["CRBCL_H1"]))
    story.append(Paragraph("• <b>Dual Authentication Credentials:</b> The system supports both standard email/password authentication (Argon2id hashing) and frontline 6-digit secure PIN authentication for rapid, touch-friendly mobile access.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Granular 114-Permission RBAC Matrix:</b> Every endpoint and database action validates discrete permissions (e.g. 'case:create', 'client:medical:read'). Fourteen canonical roles bundle standard permissions.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Case-Level Conflict Restrictions:</b> Caseworkers with personal or familial ties to clients can be restricted from specific cases, automatically blinding them from searches, dashboards, and AI queries.", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Staff Login Quarantine:</b> Historical staff records parsed from spreadsheets remain in an INACTIVE quarantine status until formally verified and activated by the System Administrator.", styles["CRBCL_Body"]))
    story.append(Spacer(1, 10))

    # Section 5
    story.append(Paragraph("5. ASK RED BEAR AI ARCHITECTURE DEEP-DIVE", styles["CRBCL_H1"]))
    story.append(Paragraph(
        "Ask Red Bear is engineered around the Gateway and Context Manager patterns to guarantee that AI assistance never bypasses data confidentiality or statutory child welfare rules:",
        styles["CRBCL_Body"]
    ))
    story.append(Paragraph("<b>Step 1 — Prompt Guard & Safety Filter:</b> Inbound prompts are scanned for prompt injection attacks or role breakout attempts.", styles["CRBCL_Body"]))
    story.append(Paragraph("<b>Step 2 — Authorization-First Context Construction:</b> Only records to which the user has explicit, un-restricted access are pulled into the retrieval context.", styles["CRBCL_Body"]))
    story.append(Paragraph("<b>Step 3 — Allowlisted Tool Execution:</b> Invokes internal deterministic tools (e.g. get_case_summary, get_overdue_goals) rather than allowing hallucinated data.", styles["CRBCL_Body"]))
    story.append(Paragraph("<b>Step 4 — Field-Level Redaction Engine:</b> Sensitive attributes (medical notes, reporter identities, financial accounts) are stripped before prompt construction.", styles["CRBCL_Body"]))
    story.append(Paragraph("<b>Step 5 — Speech Transcription as Input Aid:</b> Voice audio is transcribed in memory to the user's text box for explicit human review; audio is never auto-committed as records.", styles["CRBCL_Body"]))
    story.append(Paragraph("<b>Step 6 — Current LLM Engine State:</b> Currently uses FakeAiProvider in the backend; live production requires provisioning an enterprise API key for a Canadian-hosted sovereign LLM.", styles["CRBCL_Body"]))
    story.append(Spacer(1, 10))

    # Section 6
    story.append(Paragraph("6. MOBILE FIELD CLIENT & OFFLINE SYNCHRONIZATION", styles["CRBCL_H1"]))
    story.append(Paragraph(
        "The Flutter mobile app enables caseworkers to operate in remote reserve areas without cellular connectivity. Offline mutations are written to an encrypted local SQLite outbox table with idempotency keys. Upon reconnecting, the SyncEngine processes the queue sequentially, resolving version conflicts using server-wins timestamps.",
        styles["CRBCL_Body"]
    ))
    story.append(Spacer(1, 10))

    # Section 7
    story.append(Paragraph("7. TECHNICAL DEBT, SYSTEM RISKS & HANDOVER READINESS", styles["CRBCL_H1"]))
    story.append(Paragraph("• <b>Staging DNS & CORS Synchronization:</b> Finalize Railway backend CORS headers to match the Vercel/DigitalOcean frontend domain (https://genserver.online).", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>Third-Party Gateway Setup:</b> Provide live production credentials for Twilio Canada (SMS alerts), Geotab (telematics), and Anthropic/Azure (AI LLM).", styles["CRBCL_Body"]))
    story.append(Paragraph("• <b>End-to-End Browser UAT:</b> Run automated Playwright/Cypress test runs against live staging to record baseline browser interaction traces.", styles["CRBCL_Body"]))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Document 2 PDF successfully written to: {pdf_path}")
    return pdf_path


if __name__ == "__main__":
    build_docx()
    build_pdf()
